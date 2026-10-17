"""格式差异：比较改稿前后两份 docx，按文字对齐段落，只报告“格式”变化（文字变化交给 Word 的修订/比较功能）。
用途：导师改稿 / 合并稿之后，快速看出哪些段落的样式、字体、字号、行距、缩进、对齐被改了。"""
import difflib
from docx import Document
from . import core

PF = [("style", "样式"), ("alignment", "对齐"), ("line_spacing", "行距"), ("line_spacing_rule", "行距规则"),
      ("space_before", "段前"), ("space_after", "段后"), ("first_line_indent", "首行缩进"), ("left_indent", "左缩进")]

def _v(x):
    if x is None: return None
    if hasattr(x, "pt"): return round(x.pt, 2)
    if isinstance(x, float): return round(x, 3)
    return str(x).split(" ")[0] if hasattr(x, "name") else x

def para_sig(doc, p):
    sig = {"style": p.style.name if p.style is not None else None}
    for k, _ in PF[1:]: sig[k] = _v(core.eff_pf(p, k))
    runs = set()
    for r in core.all_runs(p):
        if not r.text.strip(): continue
        runs.add((core._eff_font_raw(doc, p, r, "eastAsia"), core._eff_font_raw(doc, p, r, "ascii"), core.eff_size(doc, p, r), bool(r.bold)))
    sig["runs"] = sorted(runs, key=str)
    return sig

def _fmt_runs(rs): return "；".join(f"{e or '-'}/{a or '-'} {s:g}pt{' 粗' if b else ''}" for e, a, s, b in rs) or "（无文字）"

def diff_format(a_path, b_path):
    da, db = Document(a_path), Document(b_path)
    pa = [p for p in da.paragraphs if p.text.strip()]; pb = [p for p in db.paragraphs if p.text.strip()]
    sm = difflib.SequenceMatcher(a=[p.text.strip() for p in pa], b=[p.text.strip() for p in pb], autojunk=False)
    out = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        pairs = list(zip(range(i1, i2), range(j1, j2))) if tag in ("equal", "replace") else []
        for i, j in pairs:
            sa, sb = para_sig(da, pa[i]), para_sig(db, pb[j])
            ch = []
            for k, name in PF:
                if sa[k] != sb[k]: ch.append(f"{name} {sa[k]} → {sb[k]}")
            if sa["runs"] != sb["runs"]: ch.append(f"字体字号 {_fmt_runs(sa['runs'])} → {_fmt_runs(sb['runs'])}")
            if ch: out.append(dict(para_a=i + 1, para_b=j + 1, text=pb[j].text.strip()[:30], text_changed=tag == "replace", changes=ch))
    return out
