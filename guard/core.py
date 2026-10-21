"""docx-format-guard 原型：有效格式计算、检查、修复。独立实现，仅依赖 python-docx/lxml。"""
import re, yaml
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ALIGN = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT,
         "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "right": WD_ALIGN_PARAGRAPH.RIGHT}
CJK = re.compile(r"[\u3400-\u9fff\u3000-\u303f\uff00-\uffef]")
FONT_ATTRS = ("eastAsia", "ascii", "hAnsi", "cs")

def all_runs(para):
    """段落内全部可见 run：直接子 run + 修订插入(w:ins/w:moveTo) + 超链接内的 run；跳过删除内容。"""
    from docx.text.run import Run
    out = []
    for child in para._p:
        if child.tag == qn("w:r"): out.append(Run(child, para))
        elif child.tag in (qn("w:ins"), qn("w:moveTo"), qn("w:hyperlink"), qn("w:smartTag")):
            out += [Run(r, para) for r in child.iter(qn("w:r"))]
    return out

def load_profile(path):
    import os
    with open(path, encoding="utf-8") as f:
        p = yaml.safe_load(f)
    p.setdefault("_base", os.path.dirname(os.path.dirname(os.path.abspath(path))))
    return p

# ---------- 有效格式（继承链：run → 段落样式 basedOn 链 → docDefaults） ----------
def _rpr_val(rpr, tag, attr="val"):
    if rpr is None: return None
    el = rpr.find(qn(f"w:{tag}"))
    if el is None: return None
    return el.get(qn(f"w:{attr}"))

def _style_chain(style):
    while style is not None:
        yield style
        style = style.base_style

def _defaults_rpr(doc):
    el = doc.styles.element.find(qn("w:docDefaults"))
    if el is None: return None
    rd = el.find(qn("w:rPrDefault"))
    return None if rd is None else rd.find(qn("w:rPr"))

FONT_ALIAS = {"SimSun": "宋体", "SimHei": "黑体", "FangSong": "仿宋", "KaiTi": "楷体", "STSong": "宋体", "Microsoft YaHei": "微软雅黑"}
def norm_font(v):
    """LibreOffice 转换会写出 '宋体;SimSun' 这类备选字体串；取首项并把英文名映射为中文名。"""
    if not v: return v
    v = v.split(";")[0].strip()
    return FONT_ALIAS.get(v, v)

def _eff_font_raw(doc, para, run, slot):
    # Word 规则：同一层级上 *Theme 属性优先于显式字体名
    for rpr in [run._r.rPr] + [st.element.rPr for st in _style_chain(para.style)] + [_defaults_rpr(doc)]:
        th = _rpr_val(rpr, "rFonts", f"{slot}Theme")
        if th: return f"theme:{th}"
        v = _rpr_val(rpr, "rFonts", slot)
        if v: return v
    return None

def eff_font(doc, para, run, slot):
    return norm_font(_eff_font_raw(doc, para, run, slot))

def eff_size(doc, para, run):
    v = _rpr_val(run._r.rPr, "sz")
    if v: return int(v) / 2
    for st in _style_chain(para.style):
        v = _rpr_val(st.element.rPr, "sz")
        if v: return int(v) / 2
    v = _rpr_val(_defaults_rpr(doc), "sz")
    return int(v) / 2 if v else 10.0

def _pf_chain(para):
    yield para.paragraph_format
    for st in _style_chain(para.style):
        yield st.paragraph_format

def eff_pf(para, name):
    for pf in _pf_chain(para):
        try: v = getattr(pf, name)
        except ValueError:            # LibreOffice 写出的 jc="start"/"end" 等 ISO 值
            jc = pf._element.pPr.find(qn("w:jc")).get(qn("w:val"))
            v = {"start": WD_ALIGN_PARAGRAPH.LEFT, "end": WD_ALIGN_PARAGRAPH.RIGHT}.get(jc)
        if v is not None: return v
    return None

# ---------- 角色识别 ----------
def classify(para, profile, default="body"):
    text = para.text.strip()
    if not text: return None
    for pat in profile.get("ignore_patterns", []):
        if re.match(pat, text): return None
    for role, r in profile["styles"].items():
        if r.get("pattern") and re.match(r["pattern"], text): return role
    return default

MATH = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
EQ_NUM = re.compile(r"[（(]\s*([A-Z]|\d+)[.\-．](\d+)\s*[)）]\s*$")

def has_math(p):
    return any(True for _ in p._p.iter(MATH + "oMath")) or any(True for _ in p._p.iter(qn("w:object")))

def heading_num_ids(doc):
    """由 guard 生成的标题多级列表（abstractNum 名为 GuardHeadings）对应的 numId。"""
    try: num = doc.part.numbering_part.element
    except Exception: return set()
    absids = {a.get(qn("w:abstractNumId")) for a in num.findall(qn("w:abstractNum"))
              if a.find(qn("w:name")) is not None and a.find(qn("w:name")).get(qn("w:val")) == "GuardHeadings"}
    return {n.get(qn("w:numId")) for n in num.findall(qn("w:num")) if n.find(qn("w:abstractNumId")).get(qn("w:val")) in absids}

def para_numpr(p):
    ppr = p._p.pPr
    np_ = None if ppr is None else ppr.find(qn("w:numPr"))
    if np_ is None: return None, None
    nid = np_.find(qn("w:numId")); il = np_.find(qn("w:ilvl"))
    return (None if nid is None else nid.get(qn("w:val"))), (0 if il is None else int(il.get(qn("w:val"))))

def is_toc_para(p):
    n = (p.style.name if p.style is not None else "").lower()
    return n.startswith("toc ") or n.startswith("contents ")

def _match_role(p, t, cands, hnum, profile):
    nid, lvl = para_numpr(p)
    for role, r in cands:                     # 1) 自动编号的标题按列表级别识别（文字里已无编号）
        if nid in hnum and r.get("num_level") == lvl: return role
    for role, r in cands:
        if r.get("has_math") == "block" and has_math(p) and EQ_NUM.search(t): return role
        if r.get("has_math") == "inline" and has_math(p) and not EQ_NUM.search(t): return role
        if r.get("after_table") and p._p.getprevious() is not None and p._p.getprevious().tag == qn("w:tbl") \
           and re.match(r["pattern"], t): return role
    for role, r in cands:
        if r.get("after_table") or r.get("has_math"): continue
        if r.get("pattern") and re.match(r["pattern"], t): return role
    for role, r in cands:
        if r.get("style_match") and p.style.name == r["style"]: return role
    return None

def roles(doc, profile, regions_out=None):
    """给每个段落分配 (段落, 角色)。支持 regions（封面/摘要/目录/正文…分区）或旧的 scope 写法。"""
    styles = profile["styles"]; hnum = heading_num_ids(doc); out = []
    ignore = [re.compile(x) for x in profile.get("ignore_patterns", [])]
    regs = profile.get("regions")
    if not regs:
        regs = [dict(name="body", start=(profile.get("scope") or {}).get("start"), default="body")]
        if profile.get("scope", {}).get("end"): regs.append(dict(name="_after", start=profile["scope"]["end"]))
        if not profile.get("scope"): regs[0]["start"] = None
    cur = None if regs[0].get("start") else regs[0]; ri = 0 if cur else -1
    default = cur.get("default") if cur else None; first = True; pending_next = False
    after = profile.get("default_role_after", {})
    for p in doc.paragraphs:
        t = p.text.strip()
        if is_toc_para(p):
            out.append((p, None))
            if regions_out is not None: regions_out[p._p] = cur["name"] if cur else None
            continue
        # 分区切换
        if pending_next and ri + 1 < len(regs):
            ri += 1; cur = regs[ri]; default = cur.get("default"); first = True
        pending_next = False
        for j in range(ri + 1, len(regs)):
            rg = regs[j]
            hit = rg.get("start") and t and re.match(rg["start"], t)
            if not hit and rg.get("start_on_role") and t:
                cands = [(k, v) for k, v in styles.items() if v.get("region", "body") == rg["name"]]
                hit = _match_role(p, t, cands, hnum, profile) == rg["start_on_role"]
            if hit: ri = j; cur = rg; default = cur.get("default"); first = True; break
        if regions_out is not None: regions_out[p._p] = cur["name"] if cur else None
        if cur is None or not t or any(x.match(t) for x in ignore):
            out.append((p, None)); continue
        cands = [(k, v) for k, v in styles.items() if v.get("region", "body") == cur["name"]]
        role = None
        if first:
            role = next((k for k, v in cands if v.get("first")), None)
        first = False
        role = role or _match_role(p, t, cands, hnum, profile) or default
        if role:
            r = styles[role]
            if r.get("resets_default"):
                default = next((x for pat, x in after.items() if re.match(pat, t)), cur.get("default"))
            if r.get("next_default"): default = r["next_default"]
        if cur.get("end_after") and re.match(cur["end_after"], t): pending_next = True
        out.append((p, role))
    return out

def _len_pt(v): return None if v is None else round(v.pt, 2)

def indent_first_pt(para, size_pt):
    """首行缩进（pt）：兼容 w:firstLine 与 w:firstLineChars（1/100 字符）。"""
    for el in [para._p.pPr] + [st.element.pPr for st in _style_chain(para.style)]:
        ind = None if el is None else el.find(qn("w:ind"))
        if ind is None: continue
        ch = ind.get(qn("w:firstLineChars"))
        if ch is not None: return int(ch) / 100 * size_pt
        fl = ind.get(qn("w:firstLine"))
        if fl is not None: return int(fl) / 20
        if ind.get(qn("w:hanging")) is not None: return -int(ind.get(qn("w:hanging"))) / 20
    return 0.0

def check_para(doc, p, r, where, add, allowed=()):
    cl = r.get("clause")
    if r.get("keep_next") and not eff_pf(p, "keep_with_next"):
        add("keep_next", where, "与下段同页", "未设置", "标题与后文分离（标题落在页底）", cl)
    al = eff_pf(p, "alignment")
    if "align" in r and al != ALIGN[r["align"]] and not (r["align"] == "left" and al is None):
        add("align", where, r["align"], str(al), "对齐", cl)
    ok, act = _line_ok(p, r)
    if not ok: add("line_spacing", where, r.get("line_exact_pt", "单倍"), act, "行距", cl)
    for k, name in (("space_before_pt", "space_before"), ("space_after_pt", "space_after")):
        if k in r:
            v = eff_pf(p, name); v = 0.0 if v is None else v.pt
            if abs(v - r[k]) > 0.5: add(name, where, r[k], round(v, 2), "段前/段后", cl)
    if "first_line_indent_chars" in r or "first_line_indent_pt" in r:
        exp = r.get("first_line_indent_pt", r.get("first_line_indent_chars", 0) * r["size_pt"])
        fi = indent_first_pt(p, r["size_pt"])
        if abs(fi - exp) > 0.6: add("indent", where, exp, round(fi, 2), "首行缩进", cl)
    if "hanging_indent_pt" in r:
        fi = indent_first_pt(p, r["size_pt"])
        if abs(-fi - r["hanging_indent_pt"]) > 0.6: add("hanging_indent", where, r["hanging_indent_pt"], round(-fi, 2), "悬挂缩进", cl)
    if "left_indent_pt" in r:
        li = eff_pf(p, "left_indent"); li = 0.0 if li is None else li.pt
        if abs(li - r["left_indent_pt"]) > 0.6: add("left_indent", where, r["left_indent_pt"], round(li, 2), "左缩进", cl)
    for ri, run in enumerate(all_runs(p)):
        if not run.text.strip(): continue
        rw = f"{where}·run{ri}"
        sz = eff_size(doc, p, run)
        if abs(sz - r["size_pt"]) > 0.25:
            add("size", rw, r["size_pt"], sz, "直接格式覆盖" if run._r.rPr is not None and run._r.rPr.find(qn("w:sz")) is not None else "样式字号", cl)
        ea = eff_font(doc, p, run, "eastAsia")
        if CJK.search(run.text) and ea != r["east_asia"]:
            add("font_east_asia", rw, r["east_asia"], ea, "中英文字体错配", cl)
        la = eff_font(doc, p, run, "ascii")
        if re.search(r"[A-Za-z0-9]", run.text) and la != r["latin"]:
            add("font_latin", rw, r["latin"], la, "中英文字体错配", cl)
        if "bold" in r:                      # 规则显式规定加粗与否时，不受 allowed_direct 豁免
            b = run.bold
            if b is None:
                for st in _style_chain(p.style):
                    if st.font.bold is not None: b = st.font.bold; break
            if bool(b) != r["bold"]: add("bold", rw, r["bold"], b, "加粗", cl)

# ---------- 检查 ----------
def _sections_for(profile, key, doc):
    start = (profile.get(key) or {}).get("from_section", 0)
    return [(i, s) for i, s in enumerate(doc.sections) if i >= start]

def _line_ok(p, r):
    ls, rule = eff_pf(p, "line_spacing"), eff_pf(p, "line_spacing_rule")
    if "line_exact_pt" in r:
        return rule == WD_LINE_SPACING.EXACTLY and ls is not None and abs(ls.pt - r["line_exact_pt"]) < 0.5, f"{rule}/{getattr(ls,'pt',ls)}"
    if r.get("line_single"):
        ok = (ls is None and rule in (None, WD_LINE_SPACING.SINGLE)) or rule == WD_LINE_SPACING.SINGLE or (isinstance(ls, float) and abs(ls - 1.0) < 0.01)
        return ok, f"{rule}/{ls}"
    return True, None

def check(path_or_doc, profile):
    doc = Document(path_or_doc) if isinstance(path_or_doc, str) else path_or_doc
    issues = []
    def add(kind, where, expect, actual, cause, rule=None):
        issues.append(dict(kind=kind, where=where, expect=expect, actual=actual, cause=cause, clause=rule))
    pg = profile["page"]; m = pg["margin_cm"]
    for i, sec in enumerate(doc.sections):
        for side in ("top", "bottom", "left", "right"):
            act = getattr(sec, f"{side}_margin").cm
            if abs(act - m[side]) > 0.05:
                add("margin", f"节{i+1}.{side}", m[side], round(act, 2), "页面设置", pg.get("clause"))
    for key, attr in (("header_distance_cm", "header_distance"), ("footer_distance_cm", "footer_distance")):
        if key in pg:
            for i, sec in _sections_for(profile, "header", doc):
                act = getattr(sec, attr).cm
                if abs(act - pg[key]) > 0.05: add(attr, f"节{i+1}", pg[key], round(act, 2), "页眉页脚距边界", pg.get("clause"))
    h = profile.get("header")
    if h:
        for i, sec in _sections_for(profile, "header", doc):
            if "text" in h:
                ht = "".join(p.text for p in sec.header.paragraphs).strip()
                if ht != h["text"]: add("header", f"节{i+1}", h["text"], ht, "页眉内容", h.get("clause"))
            if "even_text" in h:
                ht = "".join(p.text for p in sec.even_page_header.paragraphs).strip()
                if ht != h["even_text"]: add("header_even", f"节{i+1}", h["even_text"], ht, "偶数页页眉内容", h.get("clause"))
        if h.get("even_odd") and not doc.settings.odd_and_even_pages_header_footer:
            add("header_even_odd", "全文", True, False, "奇偶页不同未开启", h.get("clause"))
    allowed = set(profile.get("drift", {}).get("allowed_direct", []))
    for pi, (p, role) in enumerate(roles(doc, profile)):
        if role is None: continue
        r = profile["styles"][role]; cl = r.get("clause")
        where = f"段{pi+1}「{p.text.strip()[:12]}」"
        if r.get("require_style", True) and p.style.name != r["style"]:
            add("style", where, r["style"], p.style.name, "样式错/样式混乱", cl)
        check_para(doc, p, r, where, add, allowed)
    from . import features
    for it in features.check_all(doc, profile):
        issues.append(it)
    for it in issues: it.setdefault("fixable", True)
    return issues

# ---------- 修复 ----------
def _set_fonts(rpr_owner_el, east_asia, latin):
    rpr = rpr_owner_el.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in list(rf.attrib):
        if "Theme" in a: del rf.attrib[a]          # 去掉主题字体引用，避免覆盖
    rf.set(qn("w:eastAsia"), east_asia)
    for s in ("ascii", "hAnsi", "cs"): rf.set(qn(f"w:{s}"), latin)

def _strip_run_direct(run, allowed):
    rpr = run._r.rPr
    if rpr is None: return
    keep = {"bold": ("b", "bCs"), "italic": ("i", "iCs"), "superscript": ("vertAlign",), "subscript": ("vertAlign",)}
    keep_tags = {qn(f"w:{t}") for k in allowed for t in keep.get(k, ())}
    keep_tags |= {qn("w:rStyle"), qn("w:ins"), qn("w:del"), qn("w:rPrChange")}  # 保留修订
    for child in list(rpr):
        if child.tag not in keep_tags: rpr.remove(child)

def _strip_para_direct(p, preserve=()):
    ppr = p._p.pPr
    if ppr is None: return
    for tag in ("spacing", "ind", "jc", "rPr"):
        if tag in preserve: continue
        el = ppr.find(qn(f"w:{tag}"))
        if el is not None and not (tag == "rPr" and el.find(qn("w:ins")) is not None):
            ppr.remove(el)

class StyleResolutionError(KeyError):
    """样式在文档中解析不到（已按原生名、样式名与 styleId 两种名式查找）。"""

def _style_key(s):
    """样式名/styleId 的归一键：去非字母数字、小写。
    Word 内部名 'heading 1'、LibreOffice 转换稿的 UI 名 'Heading 1' 与
    styleId 'Heading1' 归一后同为 'heading1'，据此互相解析。"""
    return re.sub(r"[^0-9a-z]+", "", str(s).lower()) if s else ""

def find_style(doc, name):
    """按名解析样式，兼容 Word 内部名与 LO 转换稿 UI 名、styleId。
    找不到返回 None——本函数永不新建样式。"""
    if not name: return None
    try: return doc.styles[name]
    except KeyError: pass
    key = _style_key(name)
    if not key: return None
    from docx.enum.style import WD_STYLE_TYPE
    fallback = None
    for st in doc.styles:
        if _style_key(st.name) == key or _style_key(st.style_id) == key:
            if st.type == WD_STYLE_TYPE.PARAGRAPH: return st
            if fallback is None: fallback = st
    return fallback

def require_style(doc, name):
    """解析样式，解析不到显式报错——不许静默新建同名样式继续跑。"""
    st = find_style(doc, name)
    if st is None:
        raise StyleResolutionError(f"样式「{name}」在文档中解析不到（已按原生名、样式名与 styleId 两种名式查找）；拒绝静默新建同名样式")
    return st

def _ensure_style(doc, name, changes=None):
    """返回文档中已有的样式（名式不敏感，经 find_style 解析）。
    样式确实不存在时才新建——这是 fix 给论文补齐规则包自定义样式的正道，
    但新建一事必须显式记入 changes（若提供），绝不静默；样式已存在时
    更不会因名式差异造出重复样式（BUG-B1：LO 转换稿上 toc/标题样式被
    造出幽灵副本、修复写进无人使用的样式，残留 86 条且报成功）。"""
    st = find_style(doc, name)
    if st is not None: return st
    from docx.enum.style import WD_STYLE_TYPE
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    base = find_style(doc, "Normal")
    if base is not None: st.base_style = base
    if changes is not None: changes.append(f"新建样式「{name}」（文档中原本不存在）")
    return st

def _apply_style_def(doc, r, changes=None):
    st = _ensure_style(doc, r["style"], changes)
    st.font.size = Pt(r["size_pt"])
    _set_fonts(st.element, r["east_asia"], r["latin"])
    if "bold" in r: st.font.bold = r["bold"]
    pf = st.paragraph_format
    if "align" in r: pf.alignment = ALIGN[r["align"]]
    if "line_exact_pt" in r:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY; pf.line_spacing = Pt(r["line_exact_pt"])
    elif r.get("line_single"):
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if "space_before_pt" in r: pf.space_before = Pt(r["space_before_pt"])
    if "space_after_pt" in r: pf.space_after = Pt(r["space_after_pt"])
    if "first_line_indent_chars" in r or "first_line_indent_pt" in r:
        pf.first_line_indent = Pt(r.get("first_line_indent_pt", r.get("first_line_indent_chars", 0) * r["size_pt"]))
    elif "hanging_indent_pt" in r:
        pf.first_line_indent = Pt(-r["hanging_indent_pt"])
    elif r["style"] != "Normal": pf.first_line_indent = Pt(0)
    if "left_indent_pt" in r or "hanging_indent_pt" in r:
        pf.left_indent = Pt(r.get("left_indent_pt", r.get("hanging_indent_pt", 0)))
    if r.get("keep_next"): pf.keep_with_next = True
    if r.get("outline_body"):          # 前置部分标题：格式同章标题，但不进入目录
        ppr = st.element.get_or_add_pPr(); ol = ppr.find(qn("w:outlineLvl"))
        if ol is None: ol = OxmlElement("w:outlineLvl"); ppr.append(ol)
        ol.set(qn("w:val"), "9")
        np_ = ppr.find(qn("w:numPr"))
        if np_ is not None: ppr.remove(np_)
    ind = st.element.pPr.find(qn("w:ind")) if st.element.pPr is not None else None
    if ind is not None:
        for a in ("firstLineChars", "hangingChars", "leftChars", "startChars"):
            if ind.get(qn(f"w:{a}")) is not None: del ind.attrib[qn(f"w:{a}")]

def _fix_header_text(hdr, text, h, latin):
    hp = hdr.paragraphs[0]
    for extra in hdr.paragraphs[1:]:
        if not extra.text.strip(): continue
        extra._p.getparent().remove(extra._p)
    for run in list(hp.runs): run._r.getparent().remove(run._r)
    run = hp.add_run(text); run.font.size = Pt(h["size_pt"])
    _set_fonts(run._r, h["east_asia"], latin); hp.alignment = ALIGN[h["align"]]

def fix(src, dst, profile, update_toc=False):
    from . import features
    doc = Document(src)
    changes = []
    features.fix_pre(doc, profile, changes)        # 结构性修复（编号、区段、移动题注…）先于格式
    pg = profile["page"]; m = pg["margin_cm"]
    for sec in doc.sections:
        for side in ("top", "bottom", "left", "right"):
            setattr(sec, f"{side}_margin", Cm(m[side]))
    for i, sec in _sections_for(profile, "header", doc):
        if "header_distance_cm" in pg: sec.header_distance = Cm(pg["header_distance_cm"])
        if "footer_distance_cm" in pg: sec.footer_distance = Cm(pg["footer_distance_cm"])
    changes.append("页边距、页眉页脚距边界按规则重设")
    for role, r in profile["styles"].items():
        _apply_style_def(doc, r, changes)
    changes.append("样式定义按规则重写（含 eastAsia/ascii/hAnsi/cs 四槽位）")
    allowed = profile.get("drift", {}).get("allowed_direct", [])
    for p, role in roles(doc, profile):
        if role is None: continue
        r = profile["styles"][role]
        rst = require_style(doc, r["style"])
        if p.style.name != rst.name:
            changes.append(f"「{p.text[:12]}」样式 {p.style.name} → {rst.name}")
            p.style = rst
        _strip_para_direct(p, r.get("preserve_para", ()))
        role_allowed = [a for a in allowed if not (a == "bold" and "bold" in r)]
        for run in all_runs(p):
            _strip_run_direct(run, role_allowed)
    changes.append("清除范围内段落与 run 级直接格式（含修订插入的文字；保留粗斜体、上下标、修订标记）")
    h = profile.get("header")
    if h:
        latin = profile["styles"]["body"]["latin"]
        if h.get("even_odd"): doc.settings.odd_and_even_pages_header_footer = True
        for i, sec in _sections_for(profile, "header", doc):
            for key, hdr in (("text", sec.header), ("even_text", sec.even_page_header)):
                if key not in h: continue
                if "".join(p.text for p in hdr.paragraphs).strip() != h[key]:
                    hdr.is_linked_to_previous = False; _fix_header_text(hdr, h[key], h, latin)
                else:                                   # 文字正确：只规范格式，不动文字与域
                    for hp in hdr.paragraphs:
                        if not hp.text.strip(): continue
                        hp.alignment = ALIGN[h["align"]]
                        for run in hp.runs:
                            run.font.size = Pt(h["size_pt"]); _set_fonts(run._r, h["east_asia"], latin)
        changes.append("页眉按规则检查/重写（只改不合规的页眉，保留奇数页章名域）")
    used = {p.style.name for p in doc.paragraphs}
    for st in list(doc.styles):
        if st.name.startswith("Pasted") and st.name not in used:
            st.element.getparent().remove(st.element); changes.append(f"删除冗余样式 {st.name}")
    if update_toc and profile.get("toc"): profile = dict(profile, _toc_mark_dirty=False)
    features.fix_post(doc, profile, changes)
    doc.save(dst)
    if update_toc and profile.get("keep_together"):
        from .layout import fix_sig_orphans
        fix_sig_orphans(dst, profile, changes)
    if update_toc and profile.get("toc"):
        from .toc_update import update_toc_cached
        update_toc_cached(dst, profile); changes.append("目录：调用 LibreOffice 重新计算目录项与页码，写回为 TOC 域的缓存结果")
    return changes
