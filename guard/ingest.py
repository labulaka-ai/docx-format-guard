"""Derive a DRAFT rule profile from a school template .docx.

v2 reads the template's EXEMPLAR PARAGRAPHS, not just its style table:
real templates carry most of their formatting as direct formatting on the
example paragraphs, while the named style definitions can be stale (the
PKU template's Heading 1 style says 22pt Arial bold; its actual chapter
headings are 16pt SimHei, not bold). For each role, paragraphs are grouped
(by style name, with region hints for body/reference and text patterns
for figure/table captions) and the modal EFFECTIVE format — computed with
core.eff_* over the inheritance chain, the same functions check() uses —
becomes the derived rule. Style definitions are kept as a cross-check:
where they disagree with the exemplars, the disagreement is reported as a
template self-inconsistency and the exemplars win.

Headers are collected across ALL sections (v1 read only the modal-margins
section and missed headers that live in later sections).

What lives only in the written guide — region-specific role names and
semantics (the abstract/English/TOC variants), footnote numbering rules,
clause citations — still cannot be derived from a file and is listed in
the report instead of being guessed. The output is a draft: official
templates deviate from their own guides (PKU: 154 measured places), so a
review pass against the guide is required before a derived profile is
trusted.
"""
import os, re
from collections import Counter
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
import yaml
from . import core

ROLE_BY_NAME = [
    ("body", ("normal", "正文")),
    ("heading1", ("heading 1", "标题 1")),
    ("heading2", ("heading 2", "标题 2")),
    ("heading3", ("heading 3", "标题 3")),
    ("heading4", ("heading 4", "标题 4")),
    ("heading5", ("heading 5", "标题 5")),
    ("heading6", ("heading 6", "标题 6")),
    ("caption", ("caption", "题注")),
    ("reference", ("reference", "参考文献")),
    ("abstract", ("abstract", "摘要")),
]
NAME2ROLE = {a: role for role, aliases in ROLE_BY_NAME for a in aliases}

FIG_CAP = re.compile(r"^续?图\s*([A-Z]|\d+)[.\-．]\d+")
TAB_CAP = re.compile(r"^续?表\s*([A-Z]|\d+)[.\-．]\d+")
CONTENT_CHAR = re.compile(r"[㐀-鿿A-Za-z0-9]")
LATIN_CHAR = re.compile(r"[A-Za-z0-9]")

# Region anchors: light segmentation used only to keep exemplar pools clean
# (body text from the chapters, reference entries after the 参考文献 heading).
# Region NAMES of a school's rule pack are guide semantics and are not derived.
REGION_ANCHORS = [
    ("declaration", re.compile(r"学位论文原创性声明")),
    ("copyright", re.compile(r"^版权声明\s*$")),
    ("abstract_zh", re.compile(r"^摘\s*要\s*$")),
    ("abstract_en", re.compile(r"^(?i:abstract)\s*$")),
    ("toc", re.compile(r"^目\s*录\s*$")),
    ("references", re.compile(r"^参考文献\s*$")),
    ("appendix", re.compile(r"^附\s*录")),
    ("post", re.compile(r"^致\s*谢\s*$")),
    ("body", re.compile(r"^第[一二三四五六七八九十百零\d]+章")),
]

def _chain(style):
    seen, out = set(), []
    s = style
    while s is not None and id(s.element) not in seen:
        seen.add(id(s.element)); out.append(s)
        s = s.base_style
    return out

def _first(chain, xpath):
    for st in chain:
        el = st.element.find(xpath)
        if el is not None:
            return el
    return None

def _style_rule(style):
    ch = _chain(style)
    rule = {"style": style.name}
    rpr = _first(ch, qn("w:rPr"))
    if rpr is not None:
        rf = rpr.find(qn("w:rFonts"))
        if rf is not None:
            if rf.get(qn("w:eastAsia")): rule["east_asia"] = rf.get(qn("w:eastAsia"))
            if rf.get(qn("w:ascii")): rule["latin"] = rf.get(qn("w:ascii"))
        sz = rpr.find(qn("w:sz"))
        if sz is not None and sz.get(qn("w:val")):
            rule["size_pt"] = int(sz.get(qn("w:val"))) / 2
        b = rpr.find(qn("w:b"))
        if b is not None and b.get(qn("w:val")) not in ("0", "false", "off"):
            rule["bold"] = True
    ppr = _first(ch, qn("w:pPr"))
    if ppr is not None:
        sp = ppr.find(qn("w:spacing"))
        if sp is not None:
            line, rule_attr = sp.get(qn("w:line")), sp.get(qn("w:lineRule"))
            if line:
                v = int(line)
                if rule_attr == "exact": rule["line_exact_pt"] = v / 20
                elif rule_attr == "atLeast": rule["line_spacing"] = v / 20
                else: rule["line_spacing"] = round(v / 240, 3)
            if sp.get(qn("w:before")): rule["space_before_pt"] = int(sp.get(qn("w:before"))) / 20
            if sp.get(qn("w:after")): rule["space_after_pt"] = int(sp.get(qn("w:after"))) / 20
        ind = ppr.find(qn("w:ind"))
        if ind is not None:
            if ind.get(qn("w:firstLine")): rule["first_line_indent_pt"] = int(ind.get(qn("w:firstLine"))) / 20
            elif ind.get(qn("w:firstLineChars")): rule["first_line_indent_chars"] = int(ind.get(qn("w:firstLineChars"))) / 100
            if ind.get(qn("w:hanging")): rule["hanging_indent_pt"] = int(ind.get(qn("w:hanging"))) / 20
        jc = ppr.find(qn("w:jc"))
        if jc is not None and jc.get(qn("w:val")):
            rule["align"] = {"both": "justify", "distribute": "justify"}.get(jc.get(qn("w:val")), jc.get(qn("w:val")))
    return rule

# ---------- exemplar paragraphs: modal effective format per role ----------
def _modal(counter):
    return counter.most_common(1)[0][0] if counter else None

def _eff_bold(p, r):
    if r.bold is not None: return bool(r.bold)
    for st in core._style_chain(p.style):
        if st.font.bold is not None: return bool(st.font.bold)
    return False

_ALIGN_NAME = {"CENTER": "center", "LEFT": "left", "JUSTIFY": "justify", "RIGHT": "right",
               "DISTRIBUTE": "justify", "BOTH": "justify"}

def _exemplar_rule(doc, paras):
    """Modal effective format over exemplar paragraphs; returns (rule, n_paras_with_runs)."""
    ea, la, ea_all, la_all, sz, bd, styles = Counter(), Counter(), Counter(), Counter(), Counter(), Counter(), Counter()
    al, ln, sb, sa, fi = Counter(), Counter(), Counter(), Counter(), Counter()
    n = 0
    for p in paras:
        runs = [r for r in core.all_runs(p) if r.text.strip()]
        if not runs: continue
        n += 1
        if p.style is not None: styles[p.style.name] += 1
        psize = _modal(Counter(core.eff_size(doc, p, r) for r in runs))
        for r in runs:
            e, a = core.eff_font(doc, p, r, "eastAsia"), core.eff_font(doc, p, r, "ascii")
            if e: ea_all[e] += 1
            if a: la_all[a] += 1
            if core.CJK.search(r.text) and e: ea[e] += 1
            if LATIN_CHAR.search(r.text) and a: la[a] += 1
            sz[core.eff_size(doc, p, r)] += 1
            bd[_eff_bold(p, r)] += 1
        v = core.eff_pf(p, "alignment")
        if v is not None: al[str(v).split(" ")[0].upper()] += 1
        rule_e, ls = core.eff_pf(p, "line_spacing_rule"), core.eff_pf(p, "line_spacing")
        if rule_e == WD_LINE_SPACING.EXACTLY and ls is not None: ln[("exact", round(ls.pt, 2))] += 1
        elif rule_e == WD_LINE_SPACING.AT_LEAST and ls is not None: ln[("atleast", round(ls.pt, 2))] += 1
        elif rule_e == WD_LINE_SPACING.SINGLE or (rule_e is None and ls is None): ln[("single", None)] += 1
        elif isinstance(ls, float): ln[("multiple", round(ls, 3))] += 1
        for key, cnt in (("space_before", sb), ("space_after", sa)):
            v = core.eff_pf(p, key)
            cnt[0.0 if v is None else round(v.pt, 2)] += 1
        fi[round(core.indent_first_pt(p, psize or 12.0), 2)] += 1
    if not n: return None, 0
    rule = {}
    st = _modal(styles)
    if st: rule["style"] = st
    e, a = _modal(ea) or _modal(ea_all), _modal(la) or _modal(la_all)
    if e: rule["east_asia"] = e
    if a: rule["latin"] = a
    s = _modal(sz)
    if s is not None: rule["size_pt"] = s
    rule["bold"] = _modal(bd)
    av = _modal(al)
    if av in _ALIGN_NAME: rule["align"] = _ALIGN_NAME[av]
    lk = _modal(ln)
    if lk:
        if lk[0] == "exact": rule["line_exact_pt"] = lk[1]
        elif lk[0] == "single": rule["line_single"] = True
        else: rule["line_spacing"] = lk[1]
    rule["space_before_pt"] = _modal(sb); rule["space_after_pt"] = _modal(sa)
    f = _modal(fi)
    if f is not None and f <= -0.05:
        rule["hanging_indent_pt"] = round(-f, 2)
    elif f is not None and f >= 0.05:
        rule["first_line_indent_pt"] = f
        if s and round(f / s) >= 1 and abs(f / s - round(f / s)) <= 0.06:
            rule["first_line_indent_chars"] = int(round(f / s))
    return rule, n

def _is_exemplar(p):
    t = p.text.strip()
    if not t or core.is_toc_para(p): return False
    if t.startswith(("注：", "注:")): return False     # template instructions, not content exemplars
    return bool(CONTENT_CHAR.search(t))               # drops pure-placeholder paragraphs (………)

def _collect_exemplars(doc):
    paras = doc.paragraphs
    region, regs = None, []
    for p in paras:
        t = p.text.strip()
        if t:
            for rn, pat in REGION_ANCHORS:
                if pat.search(t): region = rn; break
        regs.append(region)
    has_body, has_refs = "body" in regs, "references" in regs
    roles = {}
    for p, reg in zip(paras, regs):
        if not _is_exemplar(p): continue
        t = p.text.strip()
        sname = (p.style.name if p.style is not None else "") or ""
        low = sname.lower()
        if not low.startswith(("heading", "标题")):
            if FIG_CAP.match(t): roles.setdefault("figure_caption", []).append(p); continue
            if TAB_CAP.match(t): roles.setdefault("table_caption", []).append(p); continue
        role = NAME2ROLE.get(low)
        if role == "body" and reg == "references" and has_refs:
            roles.setdefault("reference", []).append(p); continue
        if role == "body" and has_body and reg != "body":
            continue                                   # front matter / references use Normal too; not body exemplars
        if role:
            roles.setdefault(role, []).append(p)
        elif reg == "references" and has_refs and not low.startswith(("heading", "标题")):
            roles.setdefault("reference", []).append(p)
    return roles

def _paper_name(w_cm, h_cm):
    for name, (pw, ph) in {"A4": (21.0, 29.7), "A3": (29.7, 42.0), "Letter": (21.59, 27.94), "Legal": (21.59, 35.56)}.items():
        if abs(w_cm - pw) < 0.3 and abs(h_cm - ph) < 0.3: return name
    return f"{w_cm:.1f}x{h_cm:.1f}cm"

def _read_header(hdr):
    for p in hdr.paragraphs:
        t = "".join(r.text for r in p.runs).strip()
        if not t: continue
        info = {"text": t}
        if p.runs:
            rpr = p.runs[0]._r.find(qn("w:rPr"))
            if rpr is not None:
                rf = rpr.find(qn("w:rFonts"))
                if rf is not None and rf.get(qn("w:eastAsia")): info["east_asia"] = rf.get(qn("w:eastAsia"))
                sz = rpr.find(qn("w:sz"))
                if sz is not None and sz.get(qn("w:val")): info["size_pt"] = int(sz.get(qn("w:val"))) / 2
        if p.alignment is not None: info["align"] = str(p.alignment).split(" ")[0].lower().replace("both", "justify")
        return info
    return None

def _headers(doc):
    """Most common non-empty odd/even header text across ALL sections."""
    odd_c, even_c, infos = Counter(), Counter(), {}
    for s in doc.sections:
        for kind, hdr, cnt in (("odd", s.header, odd_c), ("even", s.even_page_header, even_c)):
            info = _read_header(hdr)
            if info and info.get("text"):
                cnt[info["text"]] += 1
                infos.setdefault((kind, info["text"]), info)
    return _modal(odd_c), _modal(even_c), infos, odd_c, even_c

_CONFLICT_FIELDS = ("east_asia", "latin", "size_pt", "bold", "align", "line_exact_pt")

def _conflicts(role, def_rule, ex_rule):
    out = []
    for f in _CONFLICT_FIELDS:
        if f not in def_rule or f not in ex_rule: continue
        dv, ev = def_rule[f], ex_rule[f]
        if f in ("east_asia", "latin"): dv = core.norm_font(dv)
        if isinstance(dv, (int, float)) and isinstance(ev, (int, float)):
            if abs(dv - ev) <= 0.26: continue
        elif dv == ev: continue
        out.append(f"{role}.{f}: style definition {dv!r} vs exemplar paragraphs {ev!r} (exemplars win)")
    return out

def derive_profile(template_path):
    doc = Document(template_path)
    notes, derived, conflicts = [], [], []
    name = os.path.splitext(os.path.basename(template_path))[0]
    prof = {"profile": f"{name}-draft", "derived_from_template": os.path.basename(template_path),
            "source": "AUTO-DERIVED DRAFT (v2, from the template's exemplar paragraphs; style definitions "
                      "kept as cross-check). Values were read out of the template; rules that exist only in "
                      "the written guide are NOT included (see ingest report). Review against the guide "
                      "before trusting this profile."}
    secs = []
    for i, s in enumerate(doc.sections):
        pg = s._sectPr.find(qn("w:pgNumType"))
        secs.append({"i": i, "top": round(s.top_margin.cm, 2), "bottom": round(s.bottom_margin.cm, 2),
                     "left": round(s.left_margin.cm, 2), "right": round(s.right_margin.cm, 2),
                     "header_cm": round(s.header_distance.cm, 2), "footer_cm": round(s.footer_distance.cm, 2),
                     "paper": _paper_name(round(s.page_width.cm, 2), round(s.page_height.cm, 2)),
                     "pgfmt": (pg.get(qn("w:fmt")) if pg is not None else None)})
    main = max(secs, key=lambda d: sum(1 for o in secs if (o["top"], o["bottom"], o["left"], o["right"]) == (d["top"], d["bottom"], d["left"], d["right"])))
    prof["page"] = {"margin_cm": {"top": main["top"], "bottom": main["bottom"], "left": main["left"], "right": main["right"]},
                    "header_distance_cm": main["header_cm"], "footer_distance_cm": main["footer_cm"], "paper": main["paper"]}
    derived.append(f"page setup (from section #{main['i']}, the most common margins across {len(secs)} sections)")
    if len({(s['top'], s['bottom'], s['left'], s['right']) for s in secs}) > 1:
        notes.append("sections disagree on margins; the most common set was used, all sections are listed in the report")
    # --- styles: exemplar paragraphs first, style definitions as fallback + cross-check ---
    styles, unmapped = {}, []
    by_lower = {st.name.lower(): st for st in doc.styles if st.type is not None and str(st.type).startswith("PARAGRAPH")}
    exemplars = _collect_exemplars(doc)
    exemplar_counts = {}
    def_for_role = {}
    for role, aliases in ROLE_BY_NAME:
        hit = next((by_lower[a] for a in aliases if a in by_lower), None)
        if hit is not None: def_for_role[role] = hit
    all_roles = list(dict.fromkeys([r for r, _ in ROLE_BY_NAME] + sorted(exemplars)))
    for role in all_roles:
        ex_rule, n = _exemplar_rule(doc, exemplars.get(role, []))
        def_rule = _style_rule(def_for_role[role]) if role in def_for_role else None
        if ex_rule is not None:
            exemplar_counts[role] = n
            if def_rule:
                conflicts += _conflicts(role, def_rule, ex_rule)
                for f in ("east_asia", "latin", "size_pt"):     # fill gaps exemplars cannot see
                    if f not in ex_rule and f in def_rule: ex_rule[f] = def_rule[f]
            styles[role] = ex_rule
            derived.append(f"style {role} <- {n} exemplar paragraphs")
        elif def_rule is not None:
            styles[role] = def_rule
            derived.append(f"style {role} <- style definition '{def_for_role[role].name}' (no exemplar paragraphs)")
    used = {def_for_role[r].name.lower() for r in def_for_role}
    used |= {(p.style.name.lower() if p.style is not None else "") for ps in exemplars.values() for p in ps}
    for nm, st in by_lower.items():
        if nm not in used: unmapped.append(st.name)
    if styles: prof["styles"] = styles
    if conflicts:
        notes.append("template self-inconsistencies (style definition vs its own exemplar paragraphs): "
                     + "; ".join(conflicts))
    # --- headers across all sections ---
    odd_text, even_text, hinfos, odd_c, even_c = _headers(doc)
    if odd_text or even_text:
        hdr = {}
        if odd_text:
            hdr.update(hinfos[("odd", odd_text)])
            derived.append(f"header text/format (odd pages, from {odd_c[odd_text]} section(s))")
        if even_text and even_text != odd_text:
            if not odd_text:
                base = dict(hinfos[("even", even_text)]); base.pop("text", None); hdr.update(base)
            hdr["even_text"] = even_text; hdr["even_odd"] = True
            derived.append(f"even-page header text (from {even_c[even_text]} section(s))")
        if not odd_text:
            notes.append("odd-page headers carry no plain text in any section (likely STYLEREF fields); "
                         "only the even-page header text was derived")
        prof["header"] = hdr
    fmts = [s["pgfmt"] for s in secs]
    not_derived = ["region semantics (which parts are front matter / body) and roman-vs-arabic assignment per region",
                   "region-specific role names (abstract / English-abstract / TOC variants) and roles with no "
                   "exemplar paragraphs in the template (e.g. heading4, equations, table notes) — the exemplar "
                   "paragraphs are read for the generic roles, but naming a school's region roles takes the guide",
                   "heading/caption numbering patterns and fixed texts (cover, declarations)",
                   "footnote/endnote numbering and format rules",
                   "reference-list semantics (citation style mixing, entry parsing)",
                   "clause citations — a derived profile has no guide clause numbers"]
    report = {"template": os.path.basename(template_path), "derived": derived,
              "sections": secs, "page_number_formats_by_section": fmts,
              "exemplars": exemplar_counts, "style_conflicts": conflicts,
              "styles_found_not_mapped": sorted(unmapped), "not_derived": not_derived, "notes": notes}
    return prof, report

def write_profile(prof, out_path):
    with open(out_path, "w", encoding="utf-8") as f:
        yaml.safe_dump(prof, f, allow_unicode=True, sort_keys=False)

def format_report(rep):
    lines = [f"ingest report for {rep['template']}", "derived:"]
    lines += [f"  + {d}" for d in rep["derived"]]
    lines.append("sections (margins cm / paper / page-number format):")
    for s in rep["sections"]:
        lines.append(f"  #{s['i']}: {s['top']}/{s['bottom']}/{s['left']}/{s['right']} {s['paper']} pgfmt={s['pgfmt']}")
    if rep.get("style_conflicts"):
        lines.append("template self-inconsistencies (style definition vs exemplar paragraphs; exemplars win):")
        lines += [f"  ! {c}" for c in rep["style_conflicts"]]
    if rep["styles_found_not_mapped"]:
        lines.append("styles present but not mapped to roles: " + ", ".join(rep["styles_found_not_mapped"][:20]))
    lines.append("NOT derived (need the written guide):")
    lines += [f"  - {d}" for d in rep["not_derived"]]
    lines += [f"note: {n}" for n in rep["notes"]]
    return "\n".join(lines)
