"""列出修订/批注；回复批注；在修订处加锚定批注。独立实现，直接操作 OOXML。"""
import random, datetime
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
CT_EXT = "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml"
RT_EXT = "http://schemas.microsoft.com/office/2011/relationships/commentsExtended"
CT_COM = "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"
REV_TAGS = {"ins": "插入", "del": "删除", "moveFrom": "移出", "moveTo": "移入",
            "rPrChange": "字符格式修订", "pPrChange": "段落格式修订", "sectPrChange": "节格式修订",
            "tblPrChange": "表格格式修订", "numberingChange": "编号修订"}
P14 = f"{{{W14}}}paraId"

def _text(el, tag="t"):
    return "".join(t.text or "" for t in el.iter(qn(f"w:{tag}")))

def _para_index(doc):
    return {p: i for i, p in enumerate(doc.element.body.iter(qn("w:p")))}

def _enclosing_p(el):
    while el is not None and el.tag != qn("w:p"): el = el.getparent()
    return el

def _part(doc, reltype):
    for rel in doc.part.rels.values():
        if rel.reltype == reltype and not rel.is_external: return rel.target_part
    return None

def _xml(part):
    """返回 part 的 lxml 根元素（兼容 XmlPart 与普通 Part）。"""
    if hasattr(part, "element"): return part.element
    if not hasattr(part, "_guard_el"): part._guard_el = parse_xml(part.blob)
    return part._guard_el

def list_revisions(path_or_doc):
    doc = Document(path_or_doc) if isinstance(path_or_doc, str) else path_or_doc
    pidx = _para_index(doc); out = []
    for tag, kind in REV_TAGS.items():
        for el in doc.element.body.iter(qn(f"w:{tag}")):
            if tag in ("ins", "del") and el.getparent().tag == qn("w:rPr"):
                kind2 = "段落标记" + kind; text = ""      # 段落标记的插入/删除
            else:
                kind2 = kind
                text = _text(el, "delText") if tag in ("del", "moveFrom") else _text(el)
            p = _enclosing_p(el); old = new = None
            if tag in ("rPrChange", "pPrChange"):
                old = [c.tag.split("}")[1] for c in el.iter() if c is not el and c.tag.split("}")[1] not in ("rPr", "pPr")]
                new = [c.tag.split("}")[1] for c in el.getparent() if c is not el and c.tag.split("}")[1] not in ("rPr", "pPrChange", "rPrChange")]
                text = _text(p) if p is not None else ""
            out.append(dict(id=el.get(qn("w:id")), type=kind2, tag=tag, author=el.get(qn("w:author")),
                            date=el.get(qn("w:date")), paragraph=pidx.get(p), text=text, old=old, new=new))
    return sorted(out, key=lambda r: (r["paragraph"] if r["paragraph"] is not None else -1, int(r["id"] or 0)))

def list_comments(path_or_doc):
    doc = Document(path_or_doc) if isinstance(path_or_doc, str) else path_or_doc
    cp = _part(doc, RT.COMMENTS)
    if cp is None: return []
    ext = _part(doc, RT_EXT); parent_of = {}
    if ext is not None:
        for e in _xml(ext).iter(f"{{{W15}}}commentEx"):
            parent_of[e.get(f"{{{W15}}}paraId")] = e.get(f"{{{W15}}}paraIdParent")
    pidx = _para_index(doc); anchors = {}
    starts = {s.get(qn("w:id")): s for s in doc.element.body.iter(qn("w:commentRangeStart"))}
    for cid, s in starts.items():
        ends = [e for e in doc.element.body.iter(qn("w:commentRangeEnd")) if e.get(qn("w:id")) == cid]
        texts, on = [], False
        for el in doc.element.body.iter():
            if el is s: on = True
            elif ends and el is ends[0]: break
            elif on and el.tag in (qn("w:t"), qn("w:delText")): texts.append(el.text or "")
        anchors[cid] = (pidx.get(_enclosing_p(s)), "".join(texts))
    para2cid, out = {}, []
    for c in _xml(cp).iter(qn("w:comment")):
        ps = c.findall(qn("w:p")); pid = ps[-1].get(P14) if ps else None
        cid = c.get(qn("w:id")); para2cid[pid] = cid
        out.append(dict(id=cid, author=c.get(qn("w:author")), date=c.get(qn("w:date")), text=_text(c),
                        para_id=pid, paragraph=anchors.get(cid, (None, ""))[0], anchor_text=anchors.get(cid, (None, ""))[1]))
    for c in out:
        c["reply_to"] = para2cid.get(parent_of.get(c["para_id"])) if parent_of.get(c["para_id"]) else None
    return out

# ---------- 写入 ----------
def _new_para_id(root_els):
    used = {p.get(P14) for r in root_els for p in r.iter() if p.get(P14)}
    while True:
        v = "%08X" % random.randint(1, 0x7FFFFFFF)
        if v not in used: return v

def _ensure_comments(doc):
    cp = _part(doc, RT.COMMENTS)
    if cp is None:
        cp = doc.part._comments_part       # python-docx 1.2：按需创建 comments part
    return cp

def _ensure_ext(doc):
    ext = _part(doc, RT_EXT)
    if ext is None:
        xml = f'<w15:commentsEx xmlns:w15="{W15}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"/>'.encode()
        ext = XmlPart(PackURI("/word/commentsExtended.xml"), CT_EXT, parse_xml(xml), doc.part.package)
        doc.part.relate_to(ext, RT_EXT)
    return ext

def _make_comment(doc, text, author, initials, para_id):
    cp = _ensure_comments(doc); root = _xml(cp)
    ids = [int(c.get(qn("w:id"))) for c in root.iter(qn("w:comment"))] + \
          [int(e.get(qn("w:id"))) for e in doc.element.body.iter(qn("w:commentRangeStart"))]
    cid = str(max(ids, default=-1) + 1)
    date = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    c = parse_xml(f'<w:comment xmlns:w="{W}" xmlns:w14="{W14}" w:id="{cid}" w:author="{author}" w:date="{date}" w:initials="{initials}">'
                  f'<w:p w14:paraId="{para_id}" w14:textId="77777777"><w:pPr><w:pStyle w:val="CommentText"/></w:pPr>'
                  f'<w:r><w:rPr><w:rStyle w:val="CommentReference"/></w:rPr><w:annotationRef/></w:r>'
                  f'<w:r><w:t xml:space="preserve"></w:t></w:r></w:p></w:comment>')
    list(c.iter(qn("w:t")))[0].text = text
    root.append(c)
    return cid

def _ensure_para_ids(doc):
    """给已有批注补 w14:paraId（某些生成器不写），回复需要它。"""
    root = _xml(_ensure_comments(doc))
    for c in root.iter(qn("w:comment")):
        ps = c.findall(qn("w:p"))
        if ps and not ps[-1].get(P14): ps[-1].set(P14, _new_para_id([root]))

def _anchor(start_before, end_after, cid):
    s = OxmlElement("w:commentRangeStart"); s.set(qn("w:id"), cid)
    e = OxmlElement("w:commentRangeEnd"); e.set(qn("w:id"), cid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); rs = OxmlElement("w:rStyle"); rs.set(qn("w:val"), "CommentReference")
    rpr.append(rs); r.append(rpr); ref = OxmlElement("w:commentReference"); ref.set(qn("w:id"), cid); r.append(ref)
    start_before.addprevious(s); end_after.addnext(e); e.addnext(r)

def reply_comment(src, dst, comment_id, text, author="Author", initials="A"):
    doc = Document(src); _ensure_para_ids(doc)
    root = _xml(_ensure_comments(doc))
    parent = next(c for c in root.iter(qn("w:comment")) if c.get(qn("w:id")) == str(comment_id))
    ppid = parent.findall(qn("w:p"))[-1].get(P14)
    ext = _xml(_ensure_ext(doc))
    if not any(e.get(f"{{{W15}}}paraId") == ppid for e in ext.iter(f"{{{W15}}}commentEx")):
        e = OxmlElement("w15:commentEx") if False else parse_xml(f'<w15:commentEx xmlns:w15="{W15}" w15:paraId="{ppid}" w15:done="0"/>')
        ext.append(e)
    new_pid = _new_para_id([root, ext]); cid = _make_comment(doc, text, author, initials, new_pid)
    ext.append(parse_xml(f'<w15:commentEx xmlns:w15="{W15}" w15:paraId="{new_pid}" w15:paraIdParent="{ppid}" w15:done="0"/>'))
    # Word 的做法：回复与父批注共用锚定范围
    body = doc.element.body
    s = next(x for x in body.iter(qn("w:commentRangeStart")) if x.get(qn("w:id")) == str(comment_id))
    e = next(x for x in body.iter(qn("w:commentRangeEnd")) if x.get(qn("w:id")) == str(comment_id))
    ns = OxmlElement("w:commentRangeStart"); ns.set(qn("w:id"), cid); s.addnext(ns)
    ne = OxmlElement("w:commentRangeEnd"); ne.set(qn("w:id"), cid)
    ref_run = e.getnext()
    anchor_after = ref_run if ref_run is not None and ref_run.find(qn("w:commentReference")) is not None else e
    anchor_after.addnext(ne)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); rs = OxmlElement("w:rStyle"); rs.set(qn("w:val"), "CommentReference")
    rpr.append(rs); r.append(rpr); ref = OxmlElement("w:commentReference"); ref.set(qn("w:id"), cid); r.append(ref); ne.addnext(r)
    _flush(doc); doc.save(dst); return cid

def comment_on_revision(src, dst, rev_id, text, author="Author", initials="A"):
    """修订本身不能回复：在修订所在位置加一条锚定批注。"""
    doc = Document(src)
    target = None
    for tag in REV_TAGS:
        for el in doc.element.body.iter(qn(f"w:{tag}")):
            if el.get(qn("w:id")) == str(rev_id): target = el; break
        if target is not None: break
    if target is None: raise KeyError(f"修订 {rev_id} 不存在")
    _ensure_para_ids(doc)
    root = _xml(_ensure_comments(doc)); ext = _xml(_ensure_ext(doc))
    pid = _new_para_id([root, ext]); cid = _make_comment(doc, text, author, initials, pid)
    ext.append(parse_xml(f'<w15:commentEx xmlns:w15="{W15}" w15:paraId="{pid}" w15:done="0"/>'))
    if target.tag in (qn("w:ins"), qn("w:del"), qn("w:moveFrom"), qn("w:moveTo")) and target.getparent().tag == qn("w:p"):
        _anchor(target, target, cid)                         # 包住整段修订内容
    else:                                                    # 格式修订/段落标记：锚定整段
        p = _enclosing_p(target)
        kids = [k for k in p if k.tag != qn("w:pPr")]
        _anchor(kids[0], kids[-1], cid)
    _flush(doc); doc.save(dst); return cid

def _flush(doc):
    for part in doc.part.package.iter_parts():
        if hasattr(part, "_guard_el"):
            from lxml import etree
            part._blob = etree.tostring(part._guard_el, xml_declaration=True, encoding="UTF-8", standalone=True)
