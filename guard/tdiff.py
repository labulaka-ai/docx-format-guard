"""Channel 2: thesis vs template, compared directly — no rule pack involved.

check() trusts a hand-written (or derived) list of rules; anything the list
misses is invisible to it. tdiff takes the opposite route: the template
document itself is the ground truth. Paragraphs are grouped by role (from
their style names), the template's modal effective format per role is the
reference, and every thesis paragraph is compared against it. The two
channels share no inputs, so a gap has to fool both to survive.

Deliberate choices:
- bold inside body text is NOT compared (emphasis is legitimate content);
  fonts and sizes are.
- where the template's own exemplars disagree, the modal value is the
  reference and the disagreement is reported as a template warning —
  templates can deviate from their own guides, so treat warnings seriously.
"""
from collections import Counter
from docx import Document
from docx.oxml.ns import qn
from . import core
from .diff import para_sig, PF
from .ingest import ROLE_BY_NAME

EXTRA_ROLES = [(f"heading{i}", (f"heading {i}", f"标题 {i}")) for i in range(4, 7)]
ALIASES = {}
for role, names in ROLE_BY_NAME + EXTRA_ROLES:
    for n in names: ALIASES[n] = role
STRUCTURAL = ("toc", "header", "footer", "page number")

def role_of(style_name):
    if style_name is None: return None
    low = style_name.lower()
    if low.startswith(STRUCTURAL): return "structural"
    return ALIASES.get(low)

def _modal(counter):
    return counter.most_common(1)[0][0] if counter else None

def _exemplars(doc):
    roles = {}
    for p in doc.paragraphs:
        if not p.text.strip() or core.is_toc_para(p): continue
        role = role_of(p.style.name if p.style is not None else None)
        if role in (None, "structural"): continue
        sig = para_sig(doc, p)
        roles.setdefault(role, []).append(sig)
    refs, warnings = {}, []
    for role, sigs in roles.items():
        ref = {}
        for k, label in PF:
            if k == "style": continue
            c = Counter(s[k] for s in sigs)
            ref[k] = _modal(c)
            if len(sigs) >= 3 and len(c) > 1 and c.most_common(1)[0][1] / len(sigs) < 0.8:
                warnings.append(f"模板自身不一致：角色 {role} 的{label}有 {len(c)} 种取值 {dict(c)}，已取众数")
        ea, la, sz = Counter(), Counter(), Counter()
        for s in sigs:
            for e, a, z, _b in s["runs"]:
                ea[e] += 1; la[a] += 1; sz[z] += 1
        ref["runs_modal"] = {"east_asia": _modal(ea), "latin": _modal(la), "size_pt": _modal(sz)}
        refs[role] = ref
    return refs, warnings

def _main_margins(doc):
    cs = Counter()
    for s in doc.sections:
        cs[(round(s.top_margin.cm, 2), round(s.bottom_margin.cm, 2),
            round(s.left_margin.cm, 2), round(s.right_margin.cm, 2))] += 1
    return cs.most_common(1)[0][0] if cs else None

def tdiff(thesis_path, template_path):
    th, tp = Document(thesis_path), Document(template_path)
    refs, warnings = _exemplars(tp)
    findings, unknown = [], []
    for i, p in enumerate(th.paragraphs):
        if not p.text.strip() or core.is_toc_para(p): continue
        sname = p.style.name if p.style is not None else None
        role = role_of(sname)
        if role == "structural": continue
        if role is None or role not in refs:
            unknown.append({"para": i, "style": sname, "text": p.text.strip()[:30],
                            "reason": "模板里没有这个角色/样式" if role is not None else "样式无法归入任何角色"})
            continue
        sig, ref = para_sig(th, p), refs[role]
        for k, label in PF:
            if k == "style": continue
            if sig[k] != ref[k]:
                findings.append({"para": i, "role": role, "dim": label,
                                 "template": ref[k], "actual": sig[k], "text": p.text.strip()[:30]})
        rm = ref["runs_modal"]
        for e, a, z, _b in sig["runs"]:
            for dim, want, got in (("中文字体", rm["east_asia"], e), ("西文字体", rm["latin"], a), ("字号", rm["size_pt"], z)):
                if got != want:
                    findings.append({"para": i, "role": role, "dim": dim,
                                     "template": want, "actual": got, "text": p.text.strip()[:30]})
    page = None
    ma, mb = _main_margins(th), _main_margins(tp)
    if ma != mb: page = {"template_margins_cm": mb, "thesis_margins_cm": ma}
    if len(th.sections) != len(tp.sections):
        warnings.append(f"节数不同：模板 {len(tp.sections)} 节、论文 {len(th.sections)} 节（页边距按众数节比对）")
    return {"findings": findings, "unknown": unknown, "template_warnings": warnings,
            "page_setup": page, "roles_compared": sorted(refs)}

def format_tdiff(res):
    lines = [f"模板直比（通道二，不读规则包）：比对角色 {', '.join(res['roles_compared']) or '无'}"]
    for f in res["findings"]:
        lines.append(f"  段{f['para']}「{f['text']}」{f['role']}·{f['dim']}：模板 {f['template']}，论文 {f['actual']}")
    for u in res["unknown"]:
        lines.append(f"  段{u['para']}「{u['text']}」样式 {u['style']}：{u['reason']}，未比对")
    if res["page_setup"]:
        lines.append(f"  页面边距（上/下/左/右 cm）：模板 {res['page_setup']['template_margins_cm']}，论文 {res['page_setup']['thesis_margins_cm']}")
    for w in res["template_warnings"]: lines.append("  警告：" + w)
    lines.append(f"结论：偏离 {len(res['findings'])} 处、未比对 {len(res['unknown'])} 段"
                 + ("（含页面设置不一致）" if res["page_setup"] else ""))
    return "\n".join(lines)
