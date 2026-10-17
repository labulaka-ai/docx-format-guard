"""北大规则包：原“暂不支持”条目逐项的 能检出 / 能修复 / 结构校验 / LibreOffice 渲染校验。
每个用例：在 golden（官方模板修复后的干净文档）上注入一类问题 → check 必须报出 → fix → 无可修复问题残留 → 结构校验 → 渲染校验。"""
import copy, pathlib, re, shutil, subprocess, pytest, pymupdf
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
from docx.text.paragraph import Paragraph
from guard.core import load_profile, check, fix, roles, all_runs, eff_size, eff_font, eff_pf, indent_first_pt
from guard import features as F, references as R
from wordxml import validate

ROOT = pathlib.Path(__file__).resolve().parents[1]
TPL = ROOT / "local-assets/pku/master-2024.docx"
pytestmark = pytest.mark.skipif(not TPL.exists(), reason="本地未下载北大官方模板（local-assets/ 不入库）")
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
HAS_LO = shutil.which("soffice") is not None

@pytest.fixture(scope="module")
def prof(): return load_profile(str(ROOT / "profiles/pku-master.yaml"))

@pytest.fixture(scope="module")
def golden(tmp_path_factory, prof):
    out = str(tmp_path_factory.mktemp("pkuf") / "golden.docx"); fix(str(TPL), out, prof, update_toc=True); return out

def kinds(path, prof): return {i["kind"] for i in check(path, prof)}
def fixable(iss): return [i for i in iss if i.get("fixable", True)]
def by_role(doc, prof, *want): return [p for p, r in roles(doc, prof) if r in want]

def render(path, tmp):
    if not HAS_LO: pytest.skip("无 LibreOffice")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), path], check=True, capture_output=True, timeout=240)
    return pymupdf.open(str(pathlib.Path(tmp) / (pathlib.Path(path).stem + ".pdf")))

def roundtrip(tmp_path, prof, mutate, expect, update_toc=False, allow_residual=()):
    """注入 → 检出 → 修复 → 无可修复残留 → 结构校验。返回 (bad, out)。"""
    bad, out = str(tmp_path / "bad.docx"), str(tmp_path / "out.docx")
    d = Document(GOLD[0]); mutate(d); d.save(bad)
    got = kinds(bad, prof)
    assert set(expect) <= got, f"未检出：{set(expect) - got}"
    fix(bad, out, prof, update_toc=update_toc)
    iss = check(out, prof)
    assert fixable(iss) == [], fixable(iss)[:3]
    left = {i["kind"] for i in iss} - {"ref_mixed_system", "ref_type_tag", "empty_filler"} - set(allow_residual)
    assert left == set(), left
    assert validate(out) == []
    return bad, out

GOLD = []
@pytest.fixture(autouse=True)
def _g(golden):
    GOLD[:] = [golden]

def page_texts(pdf): return [pg.get_text() for pg in pdf]

# ---------------- 1. 封面 ----------------
def test_cover(tmp_path, prof):
    def m(d):
        cell = d.tables[0].rows[0].cells[1].paragraphs[0]
        for r in cell.runs: r.font.size = Pt(20); F.core._set_fonts(r._r, "宋体", "Arial")
        info = d.tables[1].rows[0].cells[1].paragraphs[0]; info.add_run("张三").font.size = Pt(12)
        date = next(p for p in d.paragraphs if re.search("年.*月", p.text))
        F._set_text(date, "2026年6月")
    bad, out = roundtrip(tmp_path, prof, m, {"size", "font_east_asia", "cover_date_digits"})
    d = Document(out)
    assert any(p.text == "二〇二六年六月" for p in d.paragraphs)
    r = d.tables[0].rows[0].cells[1].paragraphs[0].runs[0]; p = d.tables[0].rows[0].cells[1].paragraphs[0]
    assert eff_size(d, p, r) == 26 and F.core.norm_font(eff_font(d, p, r, "eastAsia")) == "黑体"
    p = d.tables[1].rows[0].cells[1].paragraphs[0]
    assert all(eff_size(d, p, r) == 16 and F.core.norm_font(eff_font(d, p, r, "eastAsia")) == "仿宋" for r in p.runs if r.text)
    pdf = render(out, tmp_path); assert "二〇二六年六月" in page_texts(pdf)[0] + page_texts(pdf)[1]

# ---------------- 2. 版权声明 / 原创性声明（固定内容） ----------------
def test_fixed_text(tmp_path, prof):
    def m(d):
        cp = F._region_paras(d, prof, "copyright"); F._set_text(cp[-1], cp[-1].text.replace("不得", "可以"))
        dc = [p for p in F._region_paras(d, prof, "declaration") if p.text.strip()]
        for r in dc[2].runs: r.font.size = Pt(15)
    bad, out = roundtrip(tmp_path, prof, m, {"fixed_text", "fixed_format"})
    t, o = Document(str(TPL)), Document(out)
    for reg in ("copyright", "declaration"):
        assert [p.text for p in F._region_paras(t, prof, reg) if p.text.strip()] == [p.text for p in F._region_paras(o, prof, reg) if p.text.strip()]
    render(out, tmp_path)

def test_fixed_text_without_template_is_suggestion(tmp_path, prof):
    p2 = dict(prof, fixed_text={"copyright": dict(prof["fixed_text"]["copyright"], template="local-assets/不存在.docx")})
    iss = [i for i in check(GOLD[0], p2) if i["kind"] == "fixed_unverified"]
    assert iss and not iss[0]["fixable"] and iss[0]["suggestion"]

# ---------------- 3. 中英文摘要 ----------------
def test_abstract(tmp_path, prof):
    def m(d):
        t = by_role(d, prof, "en_title")[0]
        for r in t.runs: F.core._set_fonts(r._r, "宋体", "Times New Roman"); r.font.size = Pt(14)
        kw = by_role(d, prof, "en_kw")[0]; F._set_text(kw, "key words：machine learning，graph；Data")
        zk = by_role(d, prof, "abs_zh_kw")[0]; F._set_text(zk, "关键词:深度学习;图神经网络、材料")
        body = by_role(d, prof, "en_body")[0]; F._set_text(body, body.text + "（中文括号）")
    bad, out = roundtrip(tmp_path, prof, m, {"font_latin", "size", "abs_en_kw_format", "abs_zh_kw_sep", "abs_en_fullwidth"})
    d = Document(out)
    assert by_role(d, prof, "en_kw")[0].text == "KEY WORDS: Machine learning, Graph, Data"
    assert by_role(d, prof, "abs_zh_kw")[0].text == "关键词：深度学习，图神经网络，材料"
    assert "(中文括号)" in by_role(d, prof, "en_body")[0].text
    t = by_role(d, prof, "en_title")[0]
    assert all(F.core.norm_font(eff_font(d, t, r, "ascii")) == "Arial" and eff_size(d, t, r) == 16 for r in t.runs if r.text)
    assert "KEY WORDS: Machine learning" in "".join(page_texts(render(out, tmp_path)))

def test_abstract_keyword_count_is_suggestion(tmp_path, prof):
    d = Document(GOLD[0]); F._set_text(by_role(d, prof, "abs_zh_kw")[0], "关键词：甲，乙，丙，丁，戊，己"); bad = str(tmp_path / "b.docx"); d.save(bad)
    out = str(tmp_path / "o.docx"); fix(bad, out, prof)
    it = [i for i in check(out, prof) if i["kind"] == "abs_zh_kw_count"]
    assert it and not it[0]["fixable"] and "不自动增删" in it[0]["suggestion"]

# ---------------- 4. 目录 ----------------
def test_toc(tmp_path, prof):
    def m(d):
        F._toc_fields(d)[0].text = ' TOC \\o "1-2" \\h \\z \\u '
        by_role(d, prof, "toc_heading")[0].style = d.styles["Heading 1"]
        d.styles["toc 1"].font.size = Pt(10.5); d.styles["toc 2"].paragraph_format.left_indent = Pt(0)
        h2 = by_role(d, prof, "heading2")[0]; F._set_text(h2, h2.text + "（改名）")
    bad, out = roundtrip(tmp_path, prof, m, {"toc_levels", "toc_switch_u", "toc_front_included", "size", "left_indent", "toc_stale"}, update_toc=True)
    d = Document(out); ent = F.toc_entries(d)
    assert [e[1].split("\t")[0] for e in ent] == [t for _, t in F.expected_toc(d, prof)]
    assert any("（改名）" in t for _, t in ent) and "目录" not in [t.split("\t")[0] for _, t in ent]
    assert d.settings.element.find(qn("w:updateFields")) is None      # 已用排版结果写回页码，无需 Word 再提示
    pdf = render(out, tmp_path); txt = page_texts(pdf)
    # 渲染校验：目录中每章的页码 = 该章标题在 PDF 中所在页的阿拉伯页码
    body0 = next(i for i, t in enumerate(txt) if re.search(r"^第一章\s*引言$", t, re.M) and "....." not in t)
    for lv, e in ent:
        name, num = e.split("\t")
        if lv != 1: continue
        k = next(i for i in range(body0, len(txt)) if re.search(rf"^{re.escape(re.sub(r'　', '', name))}", re.sub("　", "", txt[i]), re.M))
        assert int(num) == k - body0 + 1, (name, num, k - body0 + 1)

def test_toc_without_layout_marks_dirty(tmp_path, prof):
    d = Document(GOLD[0]); h2 = by_role(d, prof, "heading2")[0]; F._set_text(h2, "新的节标题"); bad = str(tmp_path / "b.docx"); d.save(bad)
    out = str(tmp_path / "o.docx"); fix(bad, out, prof)          # 不排版：目录文字重建，页码留空并让 Word 打开时更新
    assert "toc_stale" not in kinds(out, prof)
    assert Document(out).settings.element.find(qn("w:updateFields")) is not None

# ---------------- 5. 脚注 ----------------
def test_footnotes(tmp_path, prof):
    def m(d):
        part = F._footnote_part(d); root = F._fn_root(part)
        for f in F._real_footnotes(root):
            for pel in f.findall(qn("w:p")):
                p = Paragraph(pel, d._body); p.paragraph_format.first_line_indent = Pt(0); p.paragraph_format.left_indent = Pt(0)
                for r in p.runs: r.font.size = Pt(12)
                for r in pel.iter(qn("w:r")):
                    if r.find(qn("w:footnoteRef")) is not None: F._set_vert(r, "superscript")
        F._flush_part(part)
        for ref in d.element.body.iter(qn("w:footnoteReference")): F._set_vert(ref.getparent(), "baseline")
        fp = d.sections[5]._sectPr.find(qn("w:footnotePr")); fp.find(qn("w:numFmt")).set(qn("w:val"), "decimal")
    bad, out = roundtrip(tmp_path, prof, m, {"size", "fn_hanging", "fn_mark_super", "fn_ref_not_super", "fn_numbering"})
    d = Document(out); root = F._fn_root(F._footnote_part(d))
    for f in F._real_footnotes(root):
        p = Paragraph(f.find(qn("w:p")), d._body)
        assert abs(indent_first_pt(p, 9) + 13.5) < 0.1 and all(eff_size(d, p, r) == 9 for r in p.runs if r.text.strip())
    txt = "".join(page_texts(render(out, tmp_path)))
    assert "①" in txt and "1 本研究" not in txt          # 渲染：圆圈序号

# ---------------- 6. 表格内文字、表注、题注位置 ----------------
def _note_after_table(d, text="资料来源：作者整理。"):
    tbl = d.tables[-1]._tbl
    p = parse_xml(f'<w:p xmlns:w="{W}"><w:r><w:rPr><w:sz w:val="28"/></w:rPr><w:t>{text}</w:t></w:r></w:p>'); tbl.addnext(p); return p

def test_tables(tmp_path, prof):
    def m(d):
        tbl = d.tables[-1]
        for row in tbl.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    for r in p.runs: r.font.size = Pt(9)
        cap = by_role(d, prof, "table_caption")[0]; tbl._tbl.addnext(cap._p)              # 表题跑到表下方
        fig = by_role(d, prof, "figure_caption")[0]
        img = F._neighbor(fig._p, -1); assert list(img.iter(qn("w:drawing")))
        img.addprevious(fig._p)                                                       # 图题跑到图上方
        F._replace_in_runs(fig, "\u3000", " ")
        _note_after_table(d)
    bad, out = roundtrip(tmp_path, prof, m, {"size", "space_before", "table_caption_below", "figure_caption_above", "caption_gap"})
    d = Document(out)
    tbl = d.tables[-1]._tbl
    assert re.match(r"^表", F._ptext(F._neighbor(tbl, -1)))
    fig = by_role(d, prof, "figure_caption")[0]
    assert list(F._neighbor(fig._p, -1).iter(qn("w:drawing"))) and re.match(r"^图\d+\.\d+\u3000", fig.text)
    note = by_role(d, prof, "table_note")[0]
    assert eff_size(d, note, note.runs[0]) == 10.5 and eff_pf(note, "space_after") == Pt(12)
    cellp = d.tables[-1].rows[1].cells[1].paragraphs[0]
    assert eff_size(d, cellp, cellp.runs[0]) == 11 and eff_pf(cellp, "space_before") == Pt(3)
    txt = "".join(page_texts(render(out, tmp_path)))
    assert "资料来源：作者整理。" in txt

# ---------------- 7. 公式编号 ----------------
def _eq_para(num, body="E=mc^2", sep=" "):
    return parse_xml(f'<w:p xmlns:w="{W}" xmlns:m="{M}"><m:oMath><m:r><m:t>{body}</m:t></m:r></m:oMath>'
                     f'<w:r><w:t xml:space="preserve">{sep}({num})</w:t></w:r></w:p>')

def test_equations(tmp_path, prof):
    def m(d):
        ch2 = next(p for p, r in roles(d, prof) if r == "heading1" and "研究进展" in p.text)
        nxt = ch2._p.getnext()
        a, b = _eq_para("1.1"), _eq_para("1.2", "a^2+b^2=c^2", "\u3000\u3000")
        nxt.addnext(a); a.addnext(b)
        ref = parse_xml(f'<w:p xmlns:w="{W}"><w:r><w:t>由式（1.2）与式（1.1）可得结论。</w:t></w:r></w:p>'); b.addnext(ref)
    bad, out = roundtrip(tmp_path, prof, m, {"eq_number", "eq_number_pos", "eq_tab_right", "eq_tab_center", "line_spacing"})
    d = Document(out); eqs = by_role(d, prof, "equation")
    assert [re.search(r"\((.*)\)$", p.text).group(1) for p in eqs] == ["2.1", "2.2"]
    assert any("由式（2.2）与式（2.1）可得结论。" == p.text for p in d.paragraphs)
    assert eff_pf(eqs[0], "space_before") == Pt(6)
    txt = re.sub(r"\s", "", "".join(page_texts(render(out, tmp_path))))
    assert "(2.1)" in txt and "(2.2)" in txt

def test_equation_dangling_ref_is_suggestion(tmp_path, prof):
    d = Document(GOLD[0]); body = by_role(d, prof, "body")[0]; F._set_text(body, body.text + "见式（9.9）。")
    bad = str(tmp_path / "b.docx"); d.save(bad)
    it = [i for i in check(bad, prof) if i["kind"] == "eq_ref_dangling"]
    assert it and not it[0]["fixable"]

def test_inline_math_paragraph(tmp_path, prof):
    def m(d):
        p = by_role(d, prof, "body")[0]
        om = parse_xml(f'<m:oMath xmlns:m="{M}"><m:r><m:t>x+y</m:t></m:r></m:oMath>'); p._p.append(om)
    bad, out = roundtrip(tmp_path, prof, m, {"line_spacing", "space_before"})
    p = by_role(Document(out), prof, "inline_math")[0]
    assert eff_pf(p, "space_before") == Pt(3)

# ---------------- 8. 页面、页眉、分节页码 ----------------
def test_page_sections_headers(tmp_path, prof):
    def m(d):
        st = d.settings.element; st.remove(st.find(qn("w:mirrorMargins")))
        for s in d.sections: s.gutter = Cm(1); s.page_width = Cm(21.59); s.page_height = Cm(27.94)
        bd = d.styles["Header"].element.pPr.find(qn("w:pBdr")); bd.getparent().remove(bd)
        s5 = d.sections[5]; hp = s5.header.paragraphs[0]
        for c in list(hp._p):
            if c.tag != qn("w:pPr"): hp._p.remove(c)
        hp.add_run("固定文字")
        s5._sectPr.find(qn("w:pgNumType")).set(qn("w:fmt"), "lowerRoman")
        fp = next(p for p in s5.footer.paragraphs if F._has_field(p._p, "PAGE")[0])
        r = fp.runs[0]._r; dash = copy.deepcopy(r)
        for c in list(dash):
            if c.tag != qn("w:rPr"): dash.remove(c)
        t = OxmlElement("w:t"); t.text = "- "; dash.append(t); r.addprevious(dash)
        dec = d.sections[6]; dec.footer.is_linked_to_previous = True        # 声明页继承正文页码
    bad, out = roundtrip(tmp_path, prof, m, {"mirror_margins", "gutter", "paper", "header_border", "header_odd_styleref", "pagenum_fmt", "pagenum_decor", "pagenum_present"})
    d = Document(out)
    assert all(abs(s.page_width.cm - 21) < .05 for s in d.sections)
    pdf = render(out, tmp_path); txt = page_texts(pdf)
    lab = lambda t: t.split("\n")[1] if len(t.split("\n")) > 1 else ""
    # 渲染：摘要页 I、正文第一页页码 1 且奇数页页眉 = 章名、偶数页页眉 = 北京大学硕士学位论文、声明页无页码
    zh = next(i for i, t in enumerate(txt) if t.startswith("摘要\n"))
    assert lab(txt[zh]) == "I"
    b0 = next(i for i, t in enumerate(txt) if re.match(r"^引言\n1\n", t))
    assert txt[b0 + 1].startswith("北京大学硕士学位论文\n2\n")
    last = txt[-1]; assert not re.search(r"^\d+$", last, re.M)
    assert not re.search(r"^- \d+", "".join(txt), re.M)

def test_section_split_declaration(tmp_path, prof):
    """原始模板：声明与正文同节 → 拆节；拆后不产生空白页。"""
    out = GOLD[0]; d = Document(out)
    assert F._needs_split(d, prof) == []
    pdf = render(out, tmp_path); txt = page_texts(pdf)
    assert all(t.strip() for t in txt), "出现空白页"
    i = next(i for i, t in enumerate(txt) if "原创性声明" in t and "本人郑重声明" in t)
    assert not re.search(r"^\d+$", txt[i], re.M)

# ---------------- 9. 标题自动编号 ----------------
def test_heading_numbering(tmp_path, prof):
    def m(d):
        h = by_role(d, prof, "heading1")
        F._set_text(h[0], "第一章　" + h[0].text)                      # 自动编号 + 手写编号（重复）
        h2 = by_role(d, prof, "heading2")[0]
        h2._p.pPr.remove(h2._p.pPr.find(qn("w:numPr"))); F._set_text(h2, "1.1 " + h2.text)   # 退回手写
        ref = next(p for p in h if p.text == "参考文献"); F._set_numpr(ref, sorted(F.core.heading_num_ids(d))[0], 0)
    bad, out = roundtrip(tmp_path, prof, m, {"num_double", "num_manual", "num_unexpected"})
    d = Document(out); nids = F.core.heading_num_ids(d)
    for p, r in roles(d, prof):
        if r in ("heading1", "heading2", "heading3"):
            nid, _ = F.core.para_numpr(p)
            assert (nid in nids) != bool(re.match(prof["numbering"]["unnumbered"], p.text)), p.text
            assert not re.match(prof["numbering"]["manual_prefix"], p.text)
    txt = "".join(page_texts(render(out, tmp_path)))
    assert re.search(r"第一章\s*引言", txt) and re.search(r"1\.1\s*Ⅲ族氮化物", txt) and re.search(r"第四章\s*结论", txt)
    assert not re.search(r"第[一二三四五]章\s*参考文献", txt)

def test_numbering_definition_drift(tmp_path, prof):
    def m(d):
        num = d.part.numbering_part.element
        ab = next(a for a in num.findall(qn("w:abstractNum")) if a.find(qn("w:name")) is not None and a.find(qn("w:name")).get(qn("w:val")) == "GuardHeadings")
        lv = ab.find(qn("w:lvl")); lv.find(qn("w:lvlText")).set(qn("w:val"), "Chapter %1"); lv.find(qn("w:numFmt")).set(qn("w:val"), "decimal")
    bad, out = roundtrip(tmp_path, prof, m, {"num_definition"})
    txt = "".join(page_texts(render(out, tmp_path)))
    assert "Chapter" not in txt and re.search(r"第二章\s*研究进展", txt)

# ---------------- 10. 参考文献（GB/T 7714-2015） ----------------
REFS_BAD = [
    "[1] 王明，李华，张强，赵四.《图神经网络综述》[J]. 计算机学报，2020，43(5)：1~20.",
    "[2] Smith J, Brown K, Lee M, Chen X, et al. Deep learning for materials[J]. Nature, 2019, 571(7763): pp. 95–98",
    "[3] 刘洋, 陈晨, et al. 机器学习导论[M]. 北京: 清华大学出版社, 2018.",
    "[4] OpenAI. GPT-4 technical report. https://arxiv.org/abs/2303.08774",
    "[5] John Smith. Some book title. Oxford: Oxford University Press, 2001.",
    "[6] Zhou Y. Graph theory[X]. 2010.",
]

def _replace_refs(d, prof, refs, cites):
    rp = by_role(d, prof, "reference")
    for p in rp[1:]: p._p.getparent().remove(p._p)
    F._set_text(rp[0], refs[0]); anchor = rp[0]._p
    for t in refs[1:]:
        q = copy.deepcopy(anchor); anchor.addnext(q); anchor = q; F._set_text(Paragraph(q, d._body), t)
    body = by_role(d, prof, "body")
    for p in body:                                   # 先去掉模板正文里原有的引用标注
        for mm in list(R.CITE_RE.finditer(p.text)): F._replace_in_runs(p, mm.group(0), "")
    for p, c in zip(body, cites): F._set_text(p, p.text + c)

def test_references_sequential(tmp_path, prof):
    def m(d):
        _replace_refs(d, prof, REFS_BAD, ["已有研究[3]。", "参见[1,2]与[5]。", "另见[4][6]。"])
    bad, out = roundtrip(tmp_path, prof, m,
        {"ref_fullwidth", "ref_booktitle_marks", "ref_page_range", "ref_authors", "ref_etal_lang", "ref_order", "ref_type_tag", "ref_end_period", "ref_bad_tag", "ref_western_name"},
        allow_residual={"ref_western_name", "ref_bad_tag", "ref_cite_date", "ref_j_fields"})
    d = Document(out); rs = [p.text for p in by_role(d, prof, "reference")]
    assert rs[0] == "[1] 刘洋, 陈晨, 等. 机器学习导论[M]. 北京: 清华大学出版社, 2018."
    assert rs[1] == "[2] 王明, 李华, 张强, 等. 图神经网络综述[J]. 计算机学报, 2020, 43(5): 1-20."
    assert rs[2] == "[3] Smith J, Brown K, Lee M, et al. Deep learning for materials[J]. Nature, 2019, 571(7763): 95-98."
    assert rs[4].startswith("[5] OpenAI. GPT-4 technical report[EB/OL]. https://")
    bt = [p.text for p in by_role(d, prof, "body")]
    assert bt[0].endswith("已有研究[1]。") and bt[1].endswith("参见[2,3]与[4]。") and bt[2].endswith("另见[5][6]。")
    iss = check(out, prof)
    for k in ("ref_western_name", "ref_bad_tag", "ref_cite_date"):
        it = [i for i in iss if i["kind"] == k]; assert it and not it[0]["fixable"] and it[0]["suggestion"], k
    txt = "".join(page_texts(render(out, tmp_path)))
    assert "刘洋, 陈晨, 等." in txt

def test_references_strict_2015_requires_tags(tmp_path, prof):
    refs = ["[1] 张三. 论文题目. 北京大学学报, 2020, 56(1): 1-10.", "[2] 李四. 一本书. 北京: 北京大学出版社, 2019.", "[3] 王五. 无法判断类型的条目. 2018."]
    d = Document(GOLD[0]); _replace_refs(d, prof, refs, ["[1][2][3]"]); bad = str(tmp_path / "b.docx"); d.save(bad)
    assert "ref_type_tag" not in kinds(bad, prof)                    # 北大口径（基于 2005 版 A.4.6）：非电子文献类型标识可选
    strict = dict(prof, references=dict(prof["references"], require_type_tag=True))
    assert "ref_type_tag" in kinds(bad, strict)
    out = str(tmp_path / "o.docx"); fix(bad, out, strict)
    rs = [p.text for p in by_role(Document(out), strict, "reference")]
    assert rs[0] == "[1] 张三. 论文题目[J]. 北京大学学报, 2020, 56(1): 1-10." and rs[1] == "[2] 李四. 一本书[M]. 北京: 北京大学出版社, 2019."
    left = [i for i in check(out, strict) if i["kind"] == "ref_type_tag"]
    assert len(left) == 1 and not left[0]["fixable"] and "无法" in left[0]["suggestion"]

def test_references_author_year_sorted(tmp_path, prof):
    refs = ["Zhang Q, Li P. 2009. Asian emissions[J]. Atmos Chem Phys, 9: 5131-5153.",
            "王明. 2015. 黑炭研究[J]. 环境科学, 36(1): 1-8.",
            "Bond T C. 2004. Global inventory[J]. J Geophys Res, 109: 1-43.",
            "艾华. 2012. 气溶胶[M]. 北京: 科学出版社, 2012."]
    d = Document(GOLD[0]); _replace_refs(d, prof, refs, []); bad = str(tmp_path / "b.docx"); d.save(bad)
    assert "ref_ay_order" in kinds(bad, prof)
    out = str(tmp_path / "o.docx"); fix(bad, out, prof)
    rs = [p.text for p in by_role(Document(out), prof, "reference")]
    assert [r[:2] for r in rs] == ["艾华", "王明", "Bo", "Zh"]
    assert "ref_ay_order" not in kinds(out, prof) and validate(out) == []

def test_references_mixed_and_unparsed_are_suggestions(tmp_path, prof):
    iss = [i for i in check(str(TPL), prof) if i["kind"] == "ref_mixed_system"]
    assert iss and not iss[0]["fixable"] and "无法安全转换" in iss[0]["suggestion"]

# ---------------- 汇总：全部问题叠加 + 渲染 ----------------
def test_everything_at_once(tmp_path, prof):
    def m(d):
        F._toc_fields(d)[0].text = ' TOC \\o "1-2" '
        d.settings.element.remove(d.settings.element.find(qn("w:mirrorMargins")))
        _note_after_table(d)
        _replace_refs(d, prof, REFS_BAD[:3], ["[3]", "[1,2]"])
        ch2 = next(p for p, r in roles(d, prof) if r == "heading1" and "研究进展" in p.text)
        ch2._p.addnext(_eq_para("7.7"))
    bad, out = roundtrip(tmp_path, prof, m, {"toc_levels", "mirror_margins", "ref_order", "eq_number"}, update_toc=True,
                         allow_residual={"ref_j_fields"})
    n_out = len(render(out, tmp_path))
    (tmp_path / "g").mkdir(); n_gold = len(render(GOLD[0], tmp_path / "g"))
    assert abs(n_out - n_gold) <= 2, (n_out, n_gold)
