"""tdiff (channel 2): thesis vs template compared directly, no profile involved."""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from guard.tdiff import tdiff

def _ea(style, font):
    rf = style.element.get_or_add_rPr().get_or_add_rFonts(); rf.set(qn("w:eastAsia"), font)

def _spacing(style, line=None, rule=None, first=None, jc=None):
    ppr = style.element.get_or_add_pPr()
    if line is not None:
        sp = OxmlElement("w:spacing"); sp.set(qn("w:line"), str(line)); sp.set(qn("w:lineRule"), rule); ppr.append(sp)
    if first is not None:
        ind = OxmlElement("w:ind"); ind.set(qn("w:firstLine"), str(first)); ppr.append(ind)
    if jc:
        j = OxmlElement("w:jc"); j.set(qn("w:val"), jc); ppr.append(j)

def setup_styles(doc, odd_line=False):
    n = doc.styles["Normal"]
    n.font.name = "Times New Roman"; n.font.size = Pt(12); _ea(n, "宋体")
    _spacing(n, line=400, rule="exact", first=480, jc="both")
    h = doc.styles["Heading 1"]
    h.font.name = "Times New Roman"; h.font.size = Pt(16); h.font.bold = True; _ea(h, "黑体")
    _spacing(h, jc="center")

def build(path, kind):
    doc = Document(); setup_styles(doc)
    doc.add_paragraph("第一章 绪论", style="Heading 1")
    doc.add_paragraph("模板正文示例段落，内容不重要。")
    doc.add_paragraph("模板正文示例段落之二。")
    if kind == "template_odd":
        p = doc.add_paragraph("一行距不一致的示例。")
        sp = OxmlElement("w:spacing"); sp.set(qn("w:line"), "360"); sp.set(qn("w:lineRule"), "exact")
        p._p.get_or_add_pPr().append(sp)
    if kind == "drift":
        p = doc.add_paragraph("这一段里有个字号漂移。")
        r = p.add_run("漂移"); r.font.size = Pt(14)
        q = doc.add_paragraph("这一段被直接改成居中了。"); q.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st = doc.styles.add_style("ForeignStyle", WD_STYLE_TYPE.PARAGRAPH)
        f = doc.add_paragraph("外来样式段落。"); f.style = st
    doc.save(path)

def test_clean_thesis_matches_template(tmp_path):
    tpl, th = str(tmp_path / "tpl.docx"), str(tmp_path / "th.docx")
    build(tpl, "template"); build(th, "template")
    res = tdiff(th, tpl)
    assert res["findings"] == [] and res["unknown"] == [] and res["page_setup"] is None

def test_drift_and_foreign_style_are_reported(tmp_path):
    tpl, th = str(tmp_path / "tpl.docx"), str(tmp_path / "th.docx")
    build(tpl, "template"); build(th, "drift")
    res = tdiff(th, tpl)
    dims = {(f["dim"]) for f in res["findings"]}
    assert "字号" in dims and "对齐" in dims
    assert any(u["style"] == "ForeignStyle" for u in res["unknown"])

def test_template_inconsistency_warns(tmp_path):
    tpl, th = str(tmp_path / "tpl.docx"), str(tmp_path / "th.docx")
    build(tpl, "template_odd"); build(th, "template")
    res = tdiff(th, tpl)
    assert any("模板自身不一致" in w for w in res["template_warnings"])
