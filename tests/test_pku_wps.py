"""WPS for Linux 反馈的 4 个问题的回归测试：封面单页、无空白页、声明不进目录、目录打开即有页码。
渲染用 LibreOffice（本机 WPS 没有可用的无界面转 PDF 命令行）；环境同样缺宋体/黑体/仿宋。"""
import pathlib, re, shutil, subprocess, sys, zipfile, pytest, pymupdf
from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_LINE_SPACING
from guard.core import load_profile, check, fix, roles, eff_pf
from guard import features as F
from wordxml import validate

ROOT = pathlib.Path(__file__).resolve().parents[1]
TPL = ROOT / "local-assets/pku/master-2024.docx"
pytestmark = pytest.mark.skipif(not TPL.exists(), reason="本地未下载北大官方模板")
HAS_LO = shutil.which("soffice") is not None

@pytest.fixture(scope="module")
def prof(): return load_profile(str(ROOT / "profiles/pku-master.yaml"))

@pytest.fixture(scope="module")
def cli_out(tmp_path_factory):
    """与用户相同的命令行生成 demo 文件。"""
    out = tmp_path_factory.mktemp("cli") / "pku-fixed.docx"
    subprocess.run([sys.executable, "-m", "guard.cli", "fix", str(TPL), "-o", str(out), "--profile", str(ROOT / "profiles/pku-master.yaml")],
                   cwd=ROOT, capture_output=True, timeout=300)
    assert out.exists(); return str(out)

def render(path, outdir):
    if not HAS_LO: pytest.skip("无 LibreOffice")
    outdir = pathlib.Path(outdir); outdir.mkdir(exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), path], check=True, capture_output=True, timeout=240)
    return [pg.get_text() for pg in pymupdf.open(str(outdir / (pathlib.Path(path).stem + ".pdf")))]

def cover_pages(txt): return next(i for i, t in enumerate(txt) if "版权声明" in t)

def no_simsun():
    r = subprocess.run(["fc-match", "-f", "%{family}", "SimSun"], capture_output=True, text=True)
    return "SimSun" not in r.stdout and "宋体" not in r.stdout

# ---------- 1. 封面单页 ----------
def test_cover_single_page(tmp_path, prof, cli_out):
    txt = render(cli_out, tmp_path)
    assert cover_pages(txt) == 1, "封面溢出"
    assert "二〇" in txt[0] and "导师姓名" in txt[0]
    d = Document(cli_out)
    logo = next(p for p in d.paragraphs[:3] if list(p._p.iter(qn("w:drawing"))))
    assert eff_pf(logo, "line_spacing_rule") != WD_LINE_SPACING.EXACTLY       # 校徽所在段不能是固定行距（会裁切图片）
    title = next(p for p in d.paragraphs if p.text.strip() == "硕士研究生学位论文")
    ls = eff_pf(title, "line_spacing")
    assert ls is not None and ls.pt >= 36                                   # 36pt 题目行高不小于字号
    hp = d.sections[0].header.paragraphs[0]
    assert hp._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom")).get(qn("w:val")) == "nil"   # 封面无页眉线

def test_cover_overflow_root_cause_is_font(tmp_path, prof):
    """诊断：同一环境下未经工具处理的官方模板本身就溢出（缺宋体时替代字体行高 > 网格 15.6pt）。"""
    if not no_simsun(): pytest.skip("本机有宋体，诊断不适用")
    txt = render(str(TPL), tmp_path)
    assert cover_pages(txt) == 2
    assert "cover_line_auto" in {i["kind"] for i in check(str(TPL), prof)}

def test_cover_pinning_detected_and_fixed(tmp_path, prof, cli_out):
    assert "cover_line_auto" not in {i["kind"] for i in check(cli_out, prof)}
    d = Document(cli_out)
    for p, h in [(p, None) for el, si in F.section_map(d) if si == 0 for p in [F.P(d, el)]]:
        if not list(p._p.iter(qn("w:drawing"))):
            assert eff_pf(p, "line_spacing_rule") == WD_LINE_SPACING.EXACTLY

# ---------- 2. 无空白页 ----------
def _blank_pages(txt):
    bad = []
    for i, t in enumerate(txt):
        lines = [l for l in t.split("\n") if l.strip()]
        body = [l for l in lines if not re.fullmatch(r"[IVXLC]+|\d+", l.strip())]
        if len(body) <= 1: bad.append(i)                 # 只剩页眉（和页码）
    return bad

def test_no_blank_pages(tmp_path, prof, cli_out):
    txt = render(cli_out, tmp_path)
    assert _blank_pages(txt) == [], _blank_pages(txt)
    assert F.redundant_breaks(Document(cli_out)) == []

def test_blank_page_detected_and_fixed(tmp_path, prof, cli_out):
    assert len(F.redundant_breaks(Document(str(TPL)))) >= 1               # 模板第二章末尾连续两个分页符
    assert "blank_page" in {i["kind"] for i in check(str(TPL), prof)}
    d = Document(cli_out)
    h = next(p for p, r in roles(d, prof) if r == "heading1" and "结论" in p.text)
    prev = h._p.getprevious()
    for _ in range(2):
        r = prev.makeelement(qn("w:r"), {}); br = r.makeelement(qn("w:br"), {qn("w:type"): "page"}); r.append(br); prev.append(r)
    bad = str(tmp_path / "bad.docx"); d.save(bad)
    assert "blank_page" in {i["kind"] for i in check(bad, prof)}
    out = str(tmp_path / "out.docx"); fix(bad, out, prof)
    assert "blank_page" not in {i["kind"] for i in check(out, prof)} and validate(out) == []
    assert _blank_pages(render(out, tmp_path / "r")) == []

# ---------- 3. 声明不进目录 ----------
def _with_update_fields(src, dst):
    """模拟 WPS/Word“更新目录”：打开时更新域 → LibreOffice 按大纲重建目录。"""
    zi = zipfile.ZipFile(src); zo = zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED)
    for n in zi.namelist():
        b = zi.read(n)
        if n == "word/settings.xml" and b"updateFields" not in b:
            b = b.replace(b"</w:settings>", b'<w:updateFields w:val="true"/></w:settings>')
        zo.writestr(n, b)
    zo.close(); return dst

def test_declaration_not_in_toc(tmp_path, prof, cli_out):
    d = Document(cli_out)
    assert F.toc_leaks(d, prof) == []
    decl = next(p for p in d.paragraphs if p.text.startswith("北京大学学位论文原创性声明"))
    assert F._outline_level(d, decl) == 9
    assert all("原创性声明" not in t for _, t in F.toc_entries(d))
    upd = _with_update_fields(cli_out, str(tmp_path / "upd.docx"))
    txt = render(upd, tmp_path)
    toc = next(t for t in txt if re.search(r"^目录$", t, re.M) and "...." in t)
    assert "原创性声明" not in toc and "致谢" in toc
    assert any(t.startswith("北京大学学位论文原创性声明") for t in txt)     # 声明页页眉仍为声明标题

def test_toc_leak_detected_and_fixed(tmp_path, prof):
    issues = [i for i in check(str(TPL), prof) if i["kind"] == "toc_outline_leak"]
    assert any("原创性声明" in i["where"] for i in issues)

# ---------- 4. 目录打开即有页码 ----------
def test_toc_has_page_numbers_from_cli(tmp_path, prof, cli_out):
    d = Document(cli_out); ent = F.toc_entries(d)
    assert ent and all(re.search(r"\t\d+$", t) for _, t in ent), ent
    assert d.settings.element.find(qn("w:updateFields")) is None
    txt = render(cli_out, tmp_path)
    toc = next(t for t in txt if re.search(r"^目录$", t, re.M) and "...." in t)
    nums = re.findall(r"\.{5,}(\d+)$", toc, re.M)                  # 点线引导符后的页码
    assert nums == [e.split("\t")[1] for _, e in ent], (nums, toc)
    body0 = next(i for i, t in enumerate(txt) if re.search(r"^第一章\s*引言$", t, re.M) and "...." not in t)
    for lv, e in ent:
        name, num = e.split("\t")
        k = next(i for i in range(body0, len(txt)) if re.sub(r"\s", "", name) in re.sub(r"\s", "", txt[i]))
        assert int(num) == k - body0 + 1, (name, num)

# ---------- 5. 原创性声明：签名栏不单独成页 ----------
from guard.layout import page_summary

def _summ(path, outdir):
    render(path, outdir)
    return page_summary(str(pathlib.Path(outdir) / (pathlib.Path(path).stem + ".pdf")))

def _sig_pages_ok(rows, txt):
    for i, t in enumerate(txt):
        if re.search(r"^(论文作者签名|日期)", t, re.M):
            other = [l for l in t.split("\n") if l.strip() and not re.match(r"^(论文作者签名|导师签名|日期|北京大学硕士学位论文|北京大学学位论文原创性声明)", l.strip())]
            if len(other) < 3: return False, i + 1
    return True, None

def test_declaration_signature_not_alone(tmp_path, prof, cli_out):
    txt = render(cli_out, tmp_path)
    ok, pg = _sig_pages_ok(None, txt); assert ok, f"第{pg}页只有签名栏"
    d = Document(cli_out)
    sig = [p for p in d.paragraphs if p.text.strip().startswith(("论文作者签名", "日期："))]
    assert sig and "sig_keep" not in {i["kind"] for i in check(cli_out, prof)}
    # 两个部分各自成块：“学位论文使用授权说明”与其签名栏在同一页
    i = next(k for k, t in enumerate(txt) if "学位论文使用授权说明" in t and "本人完全了解" in t)
    assert "导师签名" in txt[i]

import pytest
@pytest.mark.parametrize("k", [0, 20, 50, 60])
def test_declaration_stress_overflow_keeps_signature(tmp_path, prof, k):
    """模拟替代字体下更高的行高：把授权说明正文加长 k 段，声明可能溢出/超过一页；修复后签名栏仍不单独成页。"""
    d = Document(str(TPL))
    p = next(p for p in d.paragraphs if p.text.startswith("本人完全了解"))
    F._set_text(p, p.text + "（测试加长）" * k)
    bad = str(tmp_path / "stress.docx"); d.save(bad)
    p2 = {k2: v for k2, v in prof.items() if k2 != "fixed_text"}     # 否则 fixed_text 会把声明还原成标准措辞
    out = str(tmp_path / "stress-fixed.docx"); ch = fix(bad, out, p2, update_toc=True)
    assert Document(out).paragraphs and any(q.text.count("测试加长") == k for q in Document(out).paragraphs if q.text.startswith("本人完全了解"))
    if k >= 50: assert any("段前分页" in c for c in ch), ch       # 超过一页时触发兜底
    txt = render(out, tmp_path)
    ok, pg = _sig_pages_ok(None, txt); assert ok, f"k={k}: 第{pg}页只有签名栏"
    assert validate(out) == []

def test_declaration_keep_detected_and_fixed(tmp_path, prof, cli_out):
    assert "sig_keep" in {i["kind"] for i in check(str(TPL), prof)}
    d = Document(cli_out)
    for p in d.paragraphs: p.paragraph_format.keep_with_next = None if not p.style.name.startswith(("Heading", "NoTOC")) else p.paragraph_format.keep_with_next
    bad = str(tmp_path / "b.docx"); d.save(bad)
    assert "sig_keep" in {i["kind"] for i in check(bad, prof)}
    out = str(tmp_path / "o.docx"); fix(bad, out, prof)
    assert "sig_keep" not in {i["kind"] for i in check(out, prof)} and validate(out) == []

def test_declaration_pinned_line_heights(prof, cli_out):
    d = Document(cli_out); regs = {}; roles(d, prof, regs)
    for p in d.paragraphs:
        if regs.get(p._p) == "declaration":
            assert eff_pf(p, "line_spacing_rule") == WD_LINE_SPACING.EXACTLY, p.text[:10]

# ---------- 6. 逐页摘要：无隐藏空白页；几乎空的页全部可归因于模板占位内容 ----------
EXPECTED_NEARLY_EMPTY = {r"^关键词", r"^KEY WORDS", r"^第四章", r"^附录A"}

def test_page_summary_no_hidden_blank(tmp_path, prof, cli_out):
    rows = _summ(cli_out, tmp_path)
    assert [r["page"] for r in rows if r["blank"]] == []
    near = [r for r in rows if r["nearly_empty"]]
    assert all(any(re.match(x, r["first"]) for x in EXPECTED_NEARLY_EMPTY) for r in near), near
    assert len(near) == len(EXPECTED_NEARLY_EMPTY)
    # 页码：摘要至目录罗马数字，正文阿拉伯数字连续，封面/版权/声明无页码
    foot = [r["footer"] for r in rows]
    assert foot[:2] == ["", ""] and foot[2:7] == ["I", "II", "III", "IV", "V"] and foot[-2:] == ["", ""]
    assert foot[7:-2] == [str(k) for k in range(1, len(foot) - 8)]

def test_nearly_empty_pages_are_template_placeholders(tmp_path, prof, cli_out):
    """删掉模板占位空段 / 给空章节写一段内容后，这些“几乎空”的页即消失 → 原因是模板示例内容，不是标题与后文分离。"""
    d = Document(cli_out); regs = {}; roles(d, prof, regs)
    for p in list(d.paragraphs):
        if regs.get(p._p) in ("abstract_zh", "abstract_en") and not p.text.strip() and not list(p._p.iter(qn("w:br"))):
            p._p.getparent().remove(p._p)
    for p, r in roles(d, prof):
        if r == "heading1" and re.match(r"^(结论|附录A)", p.text):
            nxt = p._p.getnext(); F._set_text(F.P(d, nxt), "这里是正文内容。" * 20)
    out = str(tmp_path / "filled.docx"); d.save(out)
    rows = _summ(out, tmp_path)
    assert [r for r in rows if r["nearly_empty"] or r["blank"]] == [], [r for r in rows if r["nearly_empty"]]
    # 标题均已“与下段同页”（不是 keepNext 缺失）
    assert all(eff_pf(p, "keep_with_next") for p, r in roles(Document(cli_out), prof) if r in ("heading1", "heading2", "heading3"))
    assert "empty_filler" in {i["kind"] for i in check(cli_out, prof)}

def test_heading_keep_next_detected_and_fixed(tmp_path, prof, cli_out):
    d = Document(cli_out); d.styles["Heading 2"].paragraph_format.keep_with_next = False
    bad = str(tmp_path / "b.docx"); d.save(bad)
    assert "keep_next" in {i["kind"] for i in check(bad, prof)}
    out = str(tmp_path / "o.docx"); fix(bad, out, prof)
    assert "keep_next" not in {i["kind"] for i in check(out, prof)}

# ---------- 7. 分节类型 / 奇偶页：WPS “虚拟偶数页” ----------
from guard.layout import parity_notes, _num

def _sects(path):
    x = zipfile.ZipFile(path).read("word/document.xml").decode()
    return re.findall(r"<w:sectPr\b.*?</w:sectPr>", x, re.S), zipfile.ZipFile(path).read("word/settings.xml").decode()

def test_section_types_no_oddpage(cli_out):
    sects, st = _sects(cli_out)
    for s in sects:
        t = re.search(r'<w:type w:val="(\w+)"', s)
        assert t and t.group(1) == "nextPage", "分节符必须是“下一页”，不得用 oddPage/evenPage 强制奇偶页"
        assert "<w:titlePg" not in s
        assert 'w:gutter="0"' in s
    # 指南 1.11：对称页边距、奇偶页不同 —— 合规设置，必须保留
    assert "<w:mirrorMargins" in st and "<w:evenAndOddHeaders" in st
    starts = [re.search(r"<w:pgNumType([^>]*)/>", s) for s in sects]
    restart = [(re.search(r'w:fmt="(\w+)"', m.group(1)).group(1)) for m in starts if m and 'w:start="1"' in m.group(1)]
    assert restart == ["upperRoman", "decimal"]          # 只在摘要（I）和第一章（1）重新起页码

def test_roman_parse():
    assert [_num(s) for s in ("I", "IV", "V", "IX", "XII", "12", "")] == [1, 4, 5, 9, 12, 12, None]

def test_parity_note_explains_wps_virtual_page(tmp_path, cli_out):
    rows = _summ(cli_out, tmp_path)
    notes = parity_notes(rows)
    toc = next(r["page"] for r in rows if r["header"] == "目录")
    # 目录止于奇数页码 V，第一章从 1 起：恰好一处同奇，位置在目录与第一章之间；其余相邻页码奇偶交替
    assert [(n["after_page"], n["prev"], n["next"]) for n in notes] == [(toc, "V", "1")]
    synthetic = [dict(page=1, footer="IV"), dict(page=2, footer="1")]
    assert parity_notes(synthetic) == []
