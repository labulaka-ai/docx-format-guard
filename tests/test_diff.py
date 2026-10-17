"""guard diff：改稿前后格式差异。注入漂移 → diff 报出；golden 对自身 → 空；修复后对 golden → 空。"""
import pytest
from docx import Document
from fixtures import build, CORRUPTIONS
from guard.core import fix
from guard.diff import diff_format
from guard.cli import main

PARA_KINDS = ["direct_override", "font_mismatch", "theme_font", "pasted_style", "spacing_margin", "caption"]

def test_identical_is_empty(golden):
    assert diff_format(golden, golden) == []

@pytest.mark.parametrize("name", PARA_KINDS)
def test_detects_drift(tmp_path, golden, profile, name):
    """golden = 合规稿；在合规稿上注入漂移（模拟改稿后乱掉）→ diff 必须报出。"""
    from fixtures import CORRUPTIONS
    d = Document(golden); CORRUPTIONS[name][0](d); bad = str(tmp_path / f"{name}.docx"); d.save(bad)
    assert diff_format(golden, bad), name

@pytest.mark.parametrize("name", PARA_KINDS)
def test_fixed_matches_golden(tmp_path, golden, profile, name):
    bad = build(str(tmp_path / f"{name}.docx"), [name]); out = str(tmp_path / "o.docx"); fix(bad, out, profile)
    assert diff_format(golden, out) == [], diff_format(golden, out)

def test_text_edit_with_format_change(tmp_path, golden):
    d = Document(golden); p = [p for p in d.paragraphs if p.text.startswith("实验")][0]
    p.runs[0].text = p.runs[0].text + "（导师补充）"; p.runs[0].font.size = 200000  # ~15.75pt
    out = str(tmp_path / "e.docx"); d.save(out)
    res = diff_format(golden, out)
    assert len(res) == 1 and res[0]["text_changed"] and any("字体字号" in c for c in res[0]["changes"])

def test_cli_exit_codes(tmp_path, golden, capsys):
    assert main(["diff", golden, golden]) == 0
    bad = build(str(tmp_path / "b.docx"), ["direct_override"])
    assert main(["diff", golden, bad]) == 1 and "格式变化的段落" in capsys.readouterr().out
