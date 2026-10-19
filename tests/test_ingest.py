"""ingest: derive a draft profile from a synthetic template with known values."""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from guard.ingest import derive_profile
from guard.core import load_profile, check

def _set_eastasia(style, font):
    rpr = style.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts(); rf.set(qn("w:eastAsia"), font)

def _ppr_spacing(style, line=None, line_rule=None, first_line_twips=None, jc=None):
    ppr = style.element.get_or_add_pPr()
    if line is not None:
        sp = OxmlElement("w:spacing"); sp.set(qn("w:line"), str(line)); sp.set(qn("w:lineRule"), line_rule); ppr.append(sp)
    if first_line_twips is not None:
        ind = OxmlElement("w:ind"); ind.set(qn("w:firstLine"), str(first_line_twips)); ppr.append(ind)
    if jc:
        j = OxmlElement("w:jc"); j.set(qn("w:val"), jc); ppr.append(j)

def build_template(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.5)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.5)
    n = doc.styles["Normal"]
    n.font.name = "Times New Roman"; n.font.size = Pt(12); _set_eastasia(n, "宋体")
    _ppr_spacing(n, line=400, line_rule="exact", first_line_twips=480, jc="both")
    h = doc.styles["Heading 1"]
    h.font.name = "Times New Roman"; h.font.size = Pt(16); h.font.bold = True; _set_eastasia(h, "黑体")
    _ppr_spacing(h, jc="center")
    hp = sec.header.paragraphs[0]
    run = hp.add_run("测试大学硕士学位论文"); run.font.size = Pt(10.5)
    rpr = run._r.get_or_add_rPr(); rf = rpr.get_or_add_rFonts(); rf.set(qn("w:eastAsia"), "宋体")
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:fmt"), "decimal"); sec._sectPr.append(pg)
    doc.add_paragraph("正文示例")
    doc.save(path)

def test_ingest_reads_template_values(tmp_path):
    tpl = str(tmp_path / "tpl.docx"); build_template(tpl)
    prof, rep = derive_profile(tpl)
    assert prof["page"]["margin_cm"] == {"top": 2.5, "bottom": 2.5, "left": 3.0, "right": 2.5}
    assert prof["page"]["paper"] == "A4"
    body = prof["styles"]["body"]
    assert body["east_asia"] == "宋体" and body["latin"] == "Times New Roman"
    assert body["size_pt"] == 12 and body["line_exact_pt"] == 20
    assert body["first_line_indent_pt"] == 24 and body["align"] == "justify"
    h1 = prof["styles"]["heading1"]
    assert h1["east_asia"] == "黑体" and h1["size_pt"] == 16 and h1["bold"] is True and h1["align"] == "center"
    assert prof["header"]["text"] == "测试大学硕士学位论文" and prof["header"]["size_pt"] == 10.5
    assert rep["page_number_formats_by_section"] == ["decimal"]
    assert any("NOT" in s or "guide" in s for s in rep["not_derived"])

def test_derived_profile_runs_check(tmp_path):
    tpl = str(tmp_path / "tpl.docx"); build_template(tpl)
    prof, _ = derive_profile(tpl)
    out = str(tmp_path / "draft.yaml")
    from guard.ingest import write_profile
    write_profile(prof, out)
    loaded = load_profile(out)
    demo = os.path.join(os.path.dirname(__file__), "..", "demo", "broken.docx")
    iss = check(demo, loaded)  # must not crash on a draft profile; issues are fine
    assert isinstance(iss, list)
