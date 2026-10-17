"""北京大学硕士规则包：官方模板自检 + 在官方模板上注入漂移后修复。
官方模板只在本地 local-assets/（.gitignore），缺失时整组跳过。"""
import pathlib, re, shutil, subprocess, pytest, pymupdf
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from guard.core import load_profile, check, fix, roles, all_runs, eff_size, eff_font, eff_pf, indent_first_pt
from guard import revisions as rv
from wordxml import validate

ROOT = pathlib.Path(__file__).resolve().parents[1]
TPL = ROOT / "local-assets/pku/master-2024.docx"
pytestmark = pytest.mark.skipif(not TPL.exists(), reason="本地未下载北大官方模板（local-assets/ 不入库）")

@pytest.fixture(scope="module")
def prof(): return load_profile(str(ROOT / "profiles/pku-master.yaml"))

@pytest.fixture(scope="module")
def golden(tmp_path_factory, prof):
    out = str(tmp_path_factory.mktemp("pku") / "golden.docx"); fix(str(TPL), out, prof); return out

# 官方模板自检：每条问题必须落在下列已归因类别中（新出现未归因的问题 → 测试失败）
# 第三项正则：段落类问题匹配段落文字，其余（页面、页码、脚注、目录、参考文献…）匹配 where
KNOWN = [
    ("模板与规范不一致", "line_spacing|space_before|space_after", r"^(第[一二三四五]章|参考文献|附录A|致谢|摘要|目录)",
     "章标题/前置标题沿用 Word 默认 heading 1 的段前 17/段后 16.5 磅、2.41 倍行距，规范要求单倍、24/18 磅"),
    ("模板与规范不一致", "indent", r".", "正文首行缩进 425/372 twips（≈21.25/18.6pt，按五号字的 2 字符），规范为小四 2 字符 = 24pt；个别占位段为 0"),
    ("模板与规范不一致", "size", r"^(…|Ⅲ族氮化物是|表3\.5)", "示例/占位段落沿用 Normal 样式 10.5pt（五号），规范正文小四 12pt；表题个别 run 10.5pt"),
    ("模板与规范不一致", "font_latin", r"^(自20世纪|图3\.15|表3\.5|本论文是在|关键词)", "正文/题注/关键词中的西文、数字 run 直接设成宋体，规范要求 Times New Roman"),
    ("规则口径待确认", "font_latin", r"^(1\.1  Ⅲ族|附录A)", "标题中的西文字符为黑体；规范只说标题用黑体，未明确标题内西文是否用 Times New Roman（本规则按 1.6 前总则取 TNR）"),
    ("模板与规范不一致（视觉等效）", "align", r"^(\d+\.\d+|KEY WORDS)", "节标题/KEY WORDS 未显式设对齐方式；单行时显示效果相同"),
    ("模板与规范不一致", "space_before|space_after", r"^(图3\.15|ENGLISH TITLE)", "图题段前段后为 0（规范 6/12）；英文题目段后 8 磅（规范 18）"),
    ("模板与规范不一致", "bold", r"^ABSTRACT$", "“ABSTRACT”沿用 heading 1 的加粗；规范为 Arial 小四，未要求加粗"),
    ("模板与规范不一致", "hanging_indent", r"^\[\d\]", "顺序编码示例文献悬挂缩进 349 twips≈17.45pt，附录 A 示例为 2 个汉字符（21pt）"),
    ("工具要求", "num_manual", r".", "模板标题编号为手打文字；指南 1.7.1 只规定编号形式，改为多级列表自动编号是为了让目录、页眉 STYLEREF、交叉引用随章节增删自动同步"),
    ("模板自述", "toc_switch_u|toc_front_included|toc_stale", r".", "模板目录为空占位，且模板注释自述“更新目录后删除前边三项（摘要、ABSTRACT、目录）”"),
    ("模板与规范不一致", "size|left_indent|font_latin", r"^目录样式", "toc 1/2/3 样式为 10.5pt、无缩进；规范为小四、二三级缩进 1/2 字符"),
    ("模板与规范不一致", "fn_hanging|fn_mark_super", r"^脚注", "脚注无悬挂缩进、注内序号为上标；规范为悬挂 1.5 字符、注内序号用正文字体"),
    ("模板与规范不一致", "size", r"^表格", "表格内文字 10/10.5pt；规范 11pt"),
    ("模板与规范不一致", "caption_gap", r".", "图/表序号与名称间是半角空格；规范为空一个汉字符"),
    ("模板与规范不一致", "pagenum_present|pagenum_font", r"^节", "封面页脚带 PAGE 域；页码 9pt（规范五号 10.5pt）"),
    (".doc→.docx 转换或模板缺失", "header_odd_styleref|mirror_margins", r".", "奇数页页眉为空、无“对称页边距”；可能是官方 .doc 本身如此，也可能是 LibreOffice 转换丢失，无法区分"),
    ("模板结构", "section_split", r".", "英文摘要、目录与中文摘要同节（奇数页页眉无法分别显示），原创性声明与正文同节（会带页码）"),
    ("模板与规范不一致", "font_east_asia", r"^封面·date", "日期 run 为黑体；规范三号宋体"),
    ("模板与规则口径", "style", r"^(摘要|ABSTRACT|目录|ENGLISH TITLE)$|^(Author|Directed|In environ|KEY WORDS)",
     "前置标题用 Heading 1（会进目录，模板注释自述需手动删除），英文摘要用“纯文本”样式；规则用专用样式以便排除出目录、STYLEREF 取页眉"),
    ("模板内容", "blank_page", r".", "模板第二章末尾连续两个分页符（段 98、99），产生空白页；指南只要求各部分从新页开始"),
    ("模板与规范不一致", "toc_outline_leak", r".", "摘要/ABSTRACT/目录/原创性声明标题用 Heading 1，更新目录时会被列入（模板注释自述需手动删除）"),
    ("环境兼容（非规范条款）", "cover_line_auto", r".", "封面为多倍行距+文档网格；缺宋体等字体时替代字体行高超网格，封面溢出（同环境下原模板亦溢出）"),
    ("环境兼容（非规范条款）", "sig_keep", r".", "声明签名栏未设“与下段同页”，替代字体下签名栏可能被单独挤到下一页"),
    ("模板占位内容", "empty_filler", r".", "摘要/英文摘要/空章节中的占位空段落（模板示例）"),
    ("模板自述（示例文献未核准）", "ref_mixed_system|ref_type_tag|ref_fullwidth|ref_authors", r".", "模板同时演示顺序编码制与著者-出版年制，并注明“各项著录信息未核准，仅为样式参考”"),
]

def _attribute(issue, text):
    for cat, kinds, pat, why in KNOWN:
        if re.fullmatch(kinds, issue["kind"]) and re.match(pat, text): return cat
    return None

def fixable(issues): return [i for i in issues if i.get("fixable", True)]

# 修复官方模板后仍会留下的、只能给建议的问题（内容层面，工具不代改）
RESIDUAL = {"ref_mixed_system", "ref_type_tag", "empty_filler"}

def test_template_self_check_all_attributed(prof):
    doc = Document(str(TPL)); texts = {f"段{i+1}": p.text.strip() for i, p in enumerate(doc.paragraphs)}
    issues = check(str(TPL), prof)
    unknown = []
    for it in issues:
        key = it["where"].split("「")[0]
        if _attribute(it, texts.get(key, it["where"])) is None: unknown.append(it)
    assert unknown == [], f"出现未归因的问题（规则写错或工具 bug？）：{unknown[:5]}"
    kinds = {i["kind"] for i in issues}
    assert not kinds & {"margin", "header_distance", "footer_distance", "header_even", "header_even_odd", "paper", "gutter"}, kinds
    body_style = [i for i in issues if i["kind"] == "style" and not re.match(r"^段(2\d|[5-7]\d|79)「", i["where"])]
    assert body_style == [], body_style       # 正文部分样式全部正确

def test_template_fixed_is_clean_valid_and_renders(tmp_path, prof, golden):
    iss = check(golden, prof)
    assert fixable(iss) == []
    assert {i["kind"] for i in iss} == RESIDUAL and all(i["suggestion"] for i in iss)
    assert validate(golden) == []
    # 文字只允许在“手写编号 → 自动编号”和参考文献著录修正处变化
    norm = lambda t: re.sub(r"^(第[一二三四五六七八九十]+章|\d+(\.\d+){1,3})\s*", "", t)
    a = [norm(p.text) for p in Document(str(TPL)).paragraphs if p.text.strip()]
    gd = Document(golden); refs = {p.text for p, r in roles(gd, prof) if r == "reference"}
    caps = {p.text for p, r in roles(gd, prof) if r in ("figure_caption", "table_caption")}
    b = [p.text for p in gd.paragraphs if p.text.strip()]
    assert len(a) == len(b)
    diff = [(x, y) for x, y in zip(a, b) if x != y]
    assert all(y in refs or (y in caps and x.replace(" ", "\u3000", 1) == y) for x, y in diff), diff
    if shutil.which("soffice"):
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path), golden, str(TPL)], check=True, capture_output=True, timeout=180)
        a, b = pymupdf.open(str(tmp_path / "golden.pdf")), pymupdf.open(str(tmp_path / "master-2024.pdf"))
        assert abs(len(a) - len(b)) <= 2, (len(a), len(b))       # 修格式不应让篇幅剧变

# ---------- 在官方模板（修复后的 golden）上注入漂移 ----------
def _by_role(doc, prof, *want):
    return [p for p, r in roles(doc, prof) if r in want]
def _rf(run, **kw):
    rpr = run._r.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for k, v in kw.items(): rf.set(qn(f"w:{k}"), v)

def d_direct_override(doc, prof):
    for p in _by_role(doc, prof, "body"):
        for r in all_runs(p): r.font.size = Pt(14); _rf(r, eastAsia="楷体", ascii="Arial", hAnsi="Arial")
def d_font_mismatch(doc, prof):
    for p in _by_role(doc, prof, "body", "reference"):
        for r in all_runs(p): _rf(r, eastAsia="Times New Roman", ascii="宋体", hAnsi="宋体")
def d_theme_font(doc, prof):
    for p in _by_role(doc, prof, "body"):
        for r in all_runs(p):
            rpr = r._r.get_or_add_rPr(); old = rpr.find(qn("w:rFonts"))
            if old is not None: rpr.remove(old)
            rf = OxmlElement("w:rFonts"); rf.set(qn("w:asciiTheme"), "minorHAnsi"); rf.set(qn("w:eastAsia"), "等线"); rpr.insert(0, rf)
def d_pasted_style(doc, prof):
    st = doc.styles.add_style("Pasted Normal", WD_STYLE_TYPE.PARAGRAPH); st.font.size = Pt(11)
    st.paragraph_format.line_spacing = 1.0
    for p in _by_role(doc, prof, "body"): p.style = st
    for p in _by_role(doc, prof, "heading2"): p.style = doc.styles["Normal"]; [setattr(r, "bold", True) for r in p.runs]
def d_spacing_margin(doc, prof):
    for p in _by_role(doc, prof, "body"):
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE; p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(10)
    for s in doc.sections: s.left_margin = s.right_margin = Cm(3.17); s.top_margin = Cm(2.54); s.header_distance = Cm(1.5)
def d_header(doc, prof):
    hp = doc.sections[3].even_page_header.paragraphs[0]
    for r in hp.runs: r.text = ""
    hp.add_run("北京大学博士学位论文")
def d_caption(doc, prof):
    for p in _by_role(doc, prof, "figure_caption", "table_caption"):
        p.alignment = 0; p.paragraph_format.space_before = Pt(0)
        for r in p.runs: r.font.size = Pt(12)
def d_heading(doc, prof):
    for p in _by_role(doc, prof, "heading1"):
        p.paragraph_format.space_before = Pt(0)
        for r in p.runs: r.bold = True; r.font.size = Pt(18); _rf(r, eastAsia="宋体")
def d_reference(doc, prof):
    for p in _by_role(doc, prof, "reference"):
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        for r in p.runs: r.font.size = Pt(12)

DRIFTS = {
    "direct_override": (d_direct_override, {"size", "font_east_asia", "font_latin"}),
    "font_mismatch": (d_font_mismatch, {"font_east_asia", "font_latin"}),
    "theme_font": (d_theme_font, {"font_east_asia", "font_latin"}),
    "pasted_style": (d_pasted_style, {"size", "line_spacing", "style"}),
    "spacing_margin": (d_spacing_margin, {"line_spacing", "indent", "space_after", "margin", "header_distance"}),
    "header": (d_header, {"header_even"}),
    "caption": (d_caption, {"align", "size", "space_before"}),
    "heading": (d_heading, {"bold", "size", "space_before", "font_east_asia"}),
    "reference": (d_reference, {"size", "line_spacing"}),
}

def _snapshot(path, prof):
    d = Document(path); snap = []
    for p, role in roles(d, prof):
        if role is None: continue
        r = prof["styles"][role]
        runs = [(round(eff_size(d, p, x), 2), eff_font(d, p, x, "eastAsia"), eff_font(d, p, x, "ascii")) for x in all_runs(p) if x.text.strip()]
        snap.append((p.text, role, p.style.name if r.get("require_style", True) else None, eff_pf(p, "alignment"),
                     eff_pf(p, "line_spacing_rule"), eff_pf(p, "space_before"), eff_pf(p, "space_after"),
                     round(indent_first_pt(p, r["size_pt"]), 1), tuple(runs)))
    secs = [(round(s.top_margin.cm, 2), round(s.bottom_margin.cm, 2), round(s.left_margin.cm, 2), round(s.right_margin.cm, 2)) for s in d.sections]
    return snap, secs

def _inject(src, dst, names, prof):
    d = Document(src)
    for n in names: DRIFTS[n][0](d, prof)
    d.save(dst); return dst

@pytest.mark.parametrize("name", sorted(DRIFTS))
def test_drift_detected_and_fixed(tmp_path, prof, golden, name):
    bad = _inject(golden, str(tmp_path / "bad.docx"), [name], prof)
    kinds = {i["kind"] for i in check(bad, prof)}
    assert DRIFTS[name][1] <= kinds, f"{name}: 期望 {DRIFTS[name][1]}，实际 {kinds}"
    out = str(tmp_path / "out.docx"); fix(bad, out, prof)
    assert fixable(check(out, prof)) == []
    assert _snapshot(out, prof) == _snapshot(golden, prof)        # 逐段逐 run 与 golden 一致
    assert validate(out) == []

def test_all_drifts_with_review_marks(tmp_path, prof, golden):
    """全部漂移叠加 + 导师修订与批注：修回且不丢修订批注，LibreOffice 能打开。"""
    bad = _inject(golden, str(tmp_path / "bad.docx"), sorted(DRIFTS), prof)
    d = Document(bad); body = _by_role(d, prof, "body")
    ins = OxmlElement("w:ins"); ins.set(qn("w:id"), "9001"); ins.set(qn("w:author"), "导师"); ins.set(qn("w:date"), "2026-10-16T08:00:00Z")
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "30"); rpr.append(sz); r.append(rpr)
    t = OxmlElement("w:t"); t.text = "（导师补充）"; r.append(t); ins.append(r); body[0]._p.append(ins)
    d.add_comment(body[1].runs, text="请补充数据来源", author="导师", initials="DS"); d.save(bad)
    out = str(tmp_path / "out.docx"); fix(bad, out, prof)
    assert fixable(check(out, prof)) == [] and validate(out) == []
    assert {x["id"] for x in rv.list_revisions(out)} >= {"9001"} and len(rv.list_comments(out)) == 1
    if shutil.which("soffice"):
        res = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path), out], capture_output=True, timeout=180)
        assert res.returncode == 0 and (tmp_path / "out.pdf").exists()
