"""BUG-B1 回归：LibreOffice 转换稿的样式名是 UI 式（Heading 1 / TOC 1 /
Footnote Text），产品按名解析必须兼容 styleId 与两种名式；
解析不到要显式报错，绝不静默新建同名「幽灵样式」。
（2026-10-21 复核组 C 新增；背景见 docs/test-environment.md 与 BUG-B1 报告。）"""
import os, sys, zipfile
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from guard import core
from guard.features import check_toc, fix_toc

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
# LO 转换稿里的名式：styleId -> UI 名（Word 内部名则为 heading 1 / toc 1 / caption …）
UI_NAMES = {"Heading1": "Heading 1", "TOC1": "TOC 1", "TOC2": "TOC 2",
            "TOC3": "TOC 3", "FootnoteText": "Footnote Text", "Caption": "Caption"}

def _style_el(sid, name):
    el = OxmlElement("w:style")
    el.set(qn("w:type"), "paragraph"); el.set(qn("w:styleId"), sid)
    nm = OxmlElement("w:name"); nm.set(qn("w:val"), name); el.append(nm)
    bo = OxmlElement("w:basedOn"); bo.set(qn("w:val"), "Normal"); el.append(bo)
    return el

def lo_like_doc(path):
    """造一份 LibreOffice 转换风格的文档：内置样式以 UI 名存储，
    含 TOC 域与一个 Heading 1 段落。"""
    doc = Document()
    doc.add_paragraph("第一章 引言", style="Heading 1")
    p = doc.add_paragraph()
    for kind, txt in (("begin", None), ("instr", ' TOC \\o "1-3" \\h \\z '), ("separate", None)):
        r = p.add_run()
        if kind == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
        else:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
    r = p.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r._r.append(fc)
    styles_el = doc.styles.element
    by_id = {el.get(qn("w:styleId")): el for el in styles_el.findall(qn("w:style"))}
    for sid, ui in UI_NAMES.items():
        if sid in by_id:
            by_id[sid].find(qn("w:name")).set(qn("w:val"), ui)
        else:
            styles_el.append(_style_el(sid, ui))
    doc.save(path)
    return Document(path)

def _style_ids(path):
    sx = zipfile.ZipFile(path).read("word/styles.xml").decode("utf8")
    import re
    return re.findall(r'w:styleId="([^"]+)"', sx)

def _profile():
    return core.load_profile(os.path.join(os.path.dirname(__file__), "..", "profiles", "pku-master.yaml"))

def test_find_style_resolves_ui_internal_and_styleid(tmp_path):
    doc = lo_like_doc(str(tmp_path / "lo.docx"))
    for query, sid in (("toc 1", "TOC1"), ("TOC 1", "TOC1"), ("toc 3", "TOC3"),
                       ("heading 1", "Heading1"), ("Heading 1", "Heading1"),
                       ("footnote text", "FootnoteText"), ("Footnote Text", "FootnoteText"),
                       ("caption", "Caption"), ("Caption", "Caption")):
        st = core.find_style(doc, query)
        assert st is not None and st.style_id == sid, query
    assert core.find_style(doc, "不存在的样式") is None

def test_require_style_raises_instead_of_silent_create(tmp_path):
    doc = lo_like_doc(str(tmp_path / "lo.docx"))
    assert core.require_style(doc, "toc 1").style_id == "TOC1"
    with pytest.raises(core.StyleResolutionError):
        core.require_style(doc, "不存在的样式")

def test_ensure_style_never_ghosts_existing(tmp_path):
    path = str(tmp_path / "lo.docx"); doc = lo_like_doc(path)
    n0 = len(_style_ids(path))
    st = core._ensure_style(doc, "toc 1")
    assert st.name == "TOC 1" and st.style_id == "TOC1"      # 拿到的是既有样式，不是新建副本
    doc.save(path)
    ids = _style_ids(path)
    assert len(ids) == n0 and len(set(ids)) == len(ids)      # 没有新样式、没有重复 styleId

def test_ensure_style_creation_is_recorded(tmp_path):
    doc = lo_like_doc(str(tmp_path / "lo.docx"))
    changes = []
    st = core._ensure_style(doc, "Brand New Style", changes)
    assert st.name == "Brand New Style"
    assert changes == ["新建样式「Brand New Style」（文档中原本不存在）"]

def test_check_toc_no_false_missing_on_ui_names(tmp_path):
    doc = lo_like_doc(str(tmp_path / "lo.docx"))
    out = []
    check_toc(doc, _profile(), out)
    assert not [i for i in out if i["kind"] == "toc_style_missing"], out

def test_fix_toc_no_ghost_styles(tmp_path):
    path = str(tmp_path / "lo.docx"); doc = lo_like_doc(path)
    fix_toc(doc, _profile(), [])
    doc.save(path)
    ids = _style_ids(path)
    assert len(ids) == len(set(ids)), "fix 后出现重复 styleId（幽灵样式）"
    assert not [i for i in ids if i in ("toc1", "toc2", "toc3")], ids
