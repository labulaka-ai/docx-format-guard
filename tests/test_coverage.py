"""coverage: blind spots must be loud — foreign styles, unconstrained attrs, waivers."""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from guard.core import load_profile
from guard.coverage import coverage

ROOT = os.path.join(os.path.dirname(__file__), "..")

def build_doc(path):
    doc = Document()
    doc.add_paragraph("正常正文一段。")
    st = doc.styles.add_style("ForeignStyle", WD_STYLE_TYPE.PARAGRAPH)
    p = doc.add_paragraph("外来样式段落。"); p.style = st
    doc.save(path)

def test_foreign_style_and_uncovered_attrs_are_blind_spots(tmp_path):
    d = str(tmp_path / "t.docx"); build_doc(d)
    prof = load_profile(os.path.join(ROOT, "profiles", "example-master.yaml"))
    cov = coverage(d, prof)
    assert cov["unmapped_styles"].get("ForeignStyle") == 1
    assert "bold" in cov["roles"]["body"]["uncovered_attrs"]  # example body rule names no bold
    assert cov["areas"]["page_numbers"] == "uncovered"
    assert cov["verdict"] == "partial" and cov["blind_spots"] > 0

def test_waive_moves_dimensions_out_of_blind_spots(tmp_path):
    d = str(tmp_path / "t.docx"); build_doc(d)
    prof = load_profile(os.path.join(ROOT, "profiles", "example-master.yaml"))
    before = coverage(d, prof)["blind_spots"]
    prof["waive"] = ["page_numbers", "cover_fixed_texts", "header", "style:ForeignStyle",
                     "body.bold", "body.space_before_pt", "body.space_after_pt",
                     "heading1.bold", "heading1.line_spacing", "heading1.first_line_indent",
                     "heading1.space_before_pt", "heading1.space_after_pt",
                     "heading2.bold", "heading2.line_spacing", "heading2.first_line_indent",
                     "heading2.space_before_pt", "heading2.space_after_pt",
                     "caption.bold", "caption.line_spacing", "caption.first_line_indent",
                     "caption.space_before_pt", "caption.space_after_pt"]
    cov = coverage(d, prof)
    assert cov["unmapped_styles"] == {} and cov["areas"]["page_numbers"] == "waived"
    assert cov["blind_spots"] < before

def test_pku_profile_on_demo_doc_runs(tmp_path):
    prof = load_profile(os.path.join(ROOT, "profiles", "pku-master.yaml"))
    cov = coverage(os.path.join(ROOT, "demo", "broken.docx"), prof)
    assert cov["paragraphs"] > 0 and cov["areas"]["page_numbers"] == "ruled"
