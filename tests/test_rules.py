"""每条规则：注入单一漂移 → check 必须报出对应类别；fix 后逐项断言并且剩余问题为 0。"""
import shutil, hashlib, pytest
from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from fixtures import build, CORRUPTIONS
from guard.core import check, fix, eff_font, eff_size, eff_pf, classify, CJK

def test_golden_is_clean(golden, profile):
    assert check(golden, profile) == []

@pytest.mark.parametrize("name", sorted(CORRUPTIONS))
def test_detects(tmp_path, profile, name):
    p = build(str(tmp_path / f"{name}.docx"), [name])
    kinds = {i["kind"] for i in check(p, profile)}
    assert CORRUPTIONS[name][1] <= kinds, f"{name}: 期望报出 {CORRUPTIONS[name][1]}，实际 {kinds}"

def assert_conforms(path, profile):
    d = Document(path); m = profile["page"]["margin_cm"]
    for s in d.sections:                                   # 页边距
        for side in m: assert abs(getattr(s, f"{side}_margin").cm - m[side]) < 0.05
        assert "".join(p.text for p in s.header.paragraphs).strip() == profile["header"]["text"]  # 页眉
    roles = []
    for p in d.paragraphs:
        role = classify(p, profile)
        if role is None: continue
        roles.append(role); r = profile["styles"][role]
        assert p.style.name == r["style"]                  # 标题层级/样式
        assert eff_pf(p, "alignment") == {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[r["align"]]
        if "line_exact_pt" in r:                           # 行距
            assert eff_pf(p, "line_spacing_rule") == WD_LINE_SPACING.EXACTLY
            assert abs(eff_pf(p, "line_spacing").pt - r["line_exact_pt"]) < 0.5
        for run in p.runs:
            if not run.text.strip(): continue
            assert abs(eff_size(d, p, run) - r["size_pt"]) < 0.25                 # 字号
            if CJK.search(run.text): assert eff_font(d, p, run, "eastAsia") == r["east_asia"]  # 中文字体
            assert eff_font(d, p, run, "ascii") == r["latin"]                     # 西文字体
            rf = run._r.rPr.find(qn("w:rFonts")) if run._r.rPr is not None else None
            assert rf is None or not any("Theme" in a for a in rf.attrib)         # 无主题字体残留
    assert roles.count("heading1") == 1 and roles.count("heading2") == 2 and roles.count("caption") == 1  # 题注与层级数量

@pytest.mark.parametrize("name", sorted(CORRUPTIONS))
def test_fix_single(tmp_path, profile, name):
    src = build(str(tmp_path / "in.docx"), [name]); out = str(tmp_path / "out.docx")
    fix(src, out, profile)
    assert check(out, profile) == []
    assert_conforms(out, profile)

def test_fix_all_combined_with_revision(tmp_path, profile):
    src = build(str(tmp_path / "in.docx"), sorted(CORRUPTIONS), revision=True)
    before = hashlib.sha256(open(src, "rb").read()).hexdigest()
    out = str(tmp_path / "out.docx"); fix(src, out, profile)
    assert hashlib.sha256(open(src, "rb").read()).hexdigest() == before   # 原稿不被改写
    assert_conforms(out, profile)
    ins = Document(out).element.body.findall(".//" + qn("w:ins"))
    assert len(ins) == 1 and ins[0].get(qn("w:author")) == "导师"          # 修订保留
    assert "（导师补充）" in "".join(t.text for t in ins[0].iter(qn("w:t")))

def test_text_unchanged(tmp_path, profile):
    src = build(str(tmp_path / "in.docx"), sorted(CORRUPTIONS)); out = str(tmp_path / "out.docx")
    fix(src, out, profile)
    assert [p.text for p in Document(src).paragraphs] == [p.text for p in Document(out).paragraphs]

def test_pasted_style_removed(tmp_path, profile):
    src = build(str(tmp_path / "in.docx"), ["pasted_style"]); out = str(tmp_path / "out.docx")
    fix(src, out, profile)
    assert "Pasted Normal" not in [s.name for s in Document(out).styles]
